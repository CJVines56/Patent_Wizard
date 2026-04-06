from __future__ import annotations

import argparse
import csv
import gzip
import json
import os
import shlex
import shutil
import subprocess
import sys
import time
from datetime import datetime, timedelta
from itertools import product
from pathlib import Path
from typing import Any

import requests
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from backend.app.services.download import _dataset_file_base, output_file as DOWNLOAD_DIR
from backend.app.store import (
    PARAMETER_SWEEP_RESET_COLLECTIONS,
    WEAVIATE_PQ_BIT_COMPRESSION,
    WEAVIATE_PQ_CENTROIDS,
    WEAVIATE_PQ_ENABLED,
    WEAVIATE_PQ_SEGMENTS,
    WEAVIATE_PQ_TRAINING_LIMIT,
    get_client,
    load_dataset_embeddings_from_lmdb,
    reset_weaviate_state,
    store_embeddings,
)

WEAVIATE_DISK_USAGE_SSH_TARGET = os.environ.get("WEAVIATE_DISK_USAGE_SSH_TARGET", "").strip()
WEAVIATE_DISK_USAGE_SSH_ARGS = os.environ.get("WEAVIATE_DISK_USAGE_SSH_ARGS", "").strip()


def _parse_int_list(raw: str, name: str) -> list[int]:
    out: list[int] = []
    for part in str(raw or "").split(","):
        value = part.strip()
        if not value:
            continue
        try:
            parsed = int(value)
        except ValueError as exc:
            raise SystemExit(f"{name} must be a comma-separated list of integers. Bad value: {value}") from exc
        out.append(parsed)
    if not out:
        raise SystemExit(f"{name} resolved to an empty list.")
    return out


def _parse_bool(value: str, name: str) -> bool:
    normalized = str(value or "").strip().lower()
    if normalized in {"1", "true", "yes", "on"}:
        return True
    if normalized in {"0", "false", "no", "off"}:
        return False
    raise SystemExit(f"{name} must contain boolean values. Bad value: {value}")


def _parse_bool_list(raw: str, name: str) -> list[bool]:
    out: list[bool] = []
    for part in str(raw or "").split(","):
        value = part.strip()
        if not value:
            continue
        out.append(_parse_bool(value, name))
    if not out:
        raise SystemExit(f"{name} resolved to an empty list.")
    return out


def _read_csv_rows(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8", newline="") as handle:
        return [dict(row) for row in csv.DictReader(handle)]


def _json_dump(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")


def _json_load(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def _load_saved_records_from_file(path: Path):
    opener = gzip.open if path.suffix == ".gz" else Path.open
    if opener is gzip.open:
        handle = opener(path, "rt", encoding="utf-8")
    else:
        handle = opener(path, "r", encoding="utf-8")
    with handle:
        for line in handle:
            line = line.strip()
            if not line:
                continue
            yield json.loads(line)


def _upload_checkpoint_path(run_dir: Path) -> Path:
    return run_dir / "upload_state.json"


def _existing_eval_run_files(run_dir: Path) -> list[Path]:
    runs_dir = run_dir / "eval_suite" / "runs"
    if not runs_dir.exists():
        return []
    return sorted(p for p in runs_dir.glob("*.run") if p.is_file())


def _build_upload_stats_from_dataset_manifests(dataset_manifests: list[dict[str, Any]]) -> dict[str, Any]:
    dataset_uploads: list[dict[str, Any]] = []
    total_uploaded = 0
    for manifest in dataset_manifests:
        uploaded = int(manifest.get("record_count") or 0)
        total_uploaded += uploaded
        dataset_uploads.append(
            {
                "canonical_dataset_id": str(manifest.get("canonical_dataset_id") or ""),
                "uploaded_records": uploaded,
                "storage_format": str(manifest.get("storage_format") or ""),
                "storage_path": str(manifest.get("storage_path") or ""),
            }
        )
    return {
        "total_uploaded_records": total_uploaded,
        "dataset_uploads": dataset_uploads,
    }


def _current_claim_total_count() -> int | None:
    try:
        client = get_client()
    except Exception as exc:
        print(f"[warn] Could not initialize Weaviate client for resume validation: {exc}")
        return None
    try:
        existing = set(client.collections.list_all())
        if "Claim" not in existing:
            return 0
        collection = client.collections.get("Claim")
        result = collection.aggregate.over_all(total_count=True)
        total = getattr(result, "total_count", None)
        if total is None:
            return 0
        return int(total)
    except Exception as exc:
        print(f"[warn] Could not inspect live Claim collection for resume validation: {exc}")
        return None
    finally:
        try:
            client.close()
        except Exception:
            pass


def _graphql_endpoint_base() -> str:
    host = os.environ.get("WEAVIATE_HTTP_HOST", os.environ.get("WEAVIATE_LOCAL_HOST", "127.0.0.1")).strip() or "127.0.0.1"
    port = os.environ.get("WEAVIATE_HTTP_PORT", os.environ.get("WEAVIATE_LOCAL_PORT", "8080")).strip() or "8080"
    scheme = "https" if str(os.environ.get("WEAVIATE_HTTP_SECURE", "0")).strip().lower() in {"1", "true", "yes", "on"} else "http"
    return f"{scheme}://{host}:{port}"


def _graphql_claim_total_count() -> tuple[int | None, str]:
    endpoint = os.environ.get("WEAVIATE_GRAPHQL", f"{_graphql_endpoint_base()}/v1/graphql").strip()
    query = "{ Aggregate { Claim { meta { count } } } }"
    try:
        response = requests.post(endpoint, json={"query": query}, timeout=30)
        response.raise_for_status()
        payload = response.json()
        count = payload.get("data", {}).get("Aggregate", {}).get("Claim", [{}])[0].get("meta", {}).get("count")
        if count is None:
            return None, endpoint
        return int(count), endpoint
    except Exception as exc:
        print(f"[warn] Could not inspect GraphQL Claim count at {endpoint}: {exc}")
        return None, endpoint


def _write_table(ws, start_row: int, headers: list[str], rows: list[dict[str, Any]]) -> int:
    header_fill = PatternFill(fill_type="solid", fgColor="D9E1F2")
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col_idx, value=header)
        cell.font = Font(bold=True)
        cell.fill = header_fill

    row_cursor = start_row + 1
    for row in rows:
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=row_cursor, column=col_idx, value=row.get(header, ""))
        row_cursor += 1
    return row_cursor


def _auto_fit_columns(ws, *, max_col: int, max_row: int) -> None:
    for col_idx in range(1, max_col + 1):
        width = 12
        col_letter = ws.cell(row=1, column=col_idx).column_letter
        for row_idx in range(1, max_row + 1):
            value = ws.cell(row=row_idx, column=col_idx).value
            if value is None:
                continue
            width = max(width, min(80, len(str(value)) + 2))
        ws.column_dimensions[col_letter].width = width


def _build_weekly_dates(start_date: str, num_datasets: int) -> list[str]:
    anchor = datetime.strptime(start_date, "%Y-%m-%d")
    count = max(0, int(num_datasets))
    return [
        (anchor - timedelta(days=7 * idx)).strftime("%Y-%m-%d")
        for idx in range(count)
    ]


def _format_seconds(value: Any) -> str:
    if value is None:
        return "n/a"
    try:
        return f"{float(value):.1f}"
    except (TypeError, ValueError):
        return "n/a"


def _resolve_dataset_ids(
    *,
    start_date: str,
    num_datasets: int,
    dataset_ids_raw: str,
    download_path: Path,
    dataset_product: str,
) -> list[str]:
    explicit = [part.strip() for part in str(dataset_ids_raw or "").split(",") if part.strip()]
    if explicit:
        return explicit
    dataset_ids: list[str] = []
    for input_date in _build_weekly_dates(start_date, num_datasets):
        base, _, _, _ = _dataset_file_base(input_date, download_path, dataset_product)
        dataset_ids.append(base.name)
    return dataset_ids


def _find_dataset_manifest(embeddings_root: Path, dataset_id: str) -> Path | None:
    direct = embeddings_root / dataset_id / "manifest.json"
    if direct.exists():
        return direct
    matches = sorted(p for p in embeddings_root.rglob("manifest.json") if p.parent.name == dataset_id)
    if not matches:
        return None
    if len(matches) > 1:
        joined = ", ".join(str(path) for path in matches)
        raise FileNotFoundError(
            f"Ambiguous dataset manifest for {dataset_id} under {embeddings_root}. Matches: {joined}"
        )
    return matches[0]


def _resolve_storage_path(
    manifest: dict[str, Any],
    *,
    storage_format: str,
    manifest_path: Path,
    selected_root: Path,
    dataset_id: str,
) -> Path:
    candidates: list[Path] = []
    if storage_format == "lmdb":
        raw = str(manifest.get("dataset_lmdb_path") or "").strip()
        if raw:
            candidates.append(Path(raw))
        candidates.append(manifest_path.parent / "embeddings.lmdb")
        candidates.append(selected_root / dataset_id / "embeddings.lmdb")
    else:
        raw = str(manifest.get("embeddings_path") or "").strip()
        if raw:
            candidates.append(Path(raw))
        candidates.append(manifest_path.parent / "embeddings.jsonl.gz")
        candidates.append(selected_root / dataset_id / "embeddings.jsonl.gz")
    for candidate in candidates:
        if candidate.exists():
            return candidate
    joined = ", ".join(str(path) for path in candidates)
    raise FileNotFoundError(f"Missing dataset embeddings file for {dataset_id}. Tried: {joined}")


def _load_dataset_manifests(embeddings_roots: list[Path], dataset_ids: list[str]) -> list[dict[str, Any]]:
    manifests: list[dict[str, Any]] = []
    for dataset_id in dataset_ids:
        manifest_path = None
        selected_root = None
        for embeddings_root in embeddings_roots:
            candidate = _find_dataset_manifest(embeddings_root, dataset_id)
            if candidate is not None and candidate.exists():
                manifest_path = candidate
                selected_root = embeddings_root
                break
        if manifest_path is None or selected_root is None:
            searched = ", ".join(str(root) for root in embeddings_roots)
            raise FileNotFoundError(f"Missing dataset manifest for {dataset_id}. Searched: {searched}")
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        storage_format = str(manifest.get("storage_format") or "jsonl.gz").strip().lower()
        storage_path = _resolve_storage_path(
            manifest,
            storage_format=storage_format,
            manifest_path=manifest_path,
            selected_root=selected_root,
            dataset_id=dataset_id,
        )
        manifest["storage_format"] = storage_format
        manifest["storage_path"] = str(storage_path.resolve())
        manifest["manifest_path"] = str(manifest_path.resolve())
        manifest["embeddings_root"] = str(selected_root.resolve())
        manifests.append(manifest)
    return manifests


def _human_size_to_bytes(raw: str) -> int | None:
    value = str(raw or "").strip()
    if not value:
        return None
    suffix = value[-1].upper()
    multipliers = {
        "B": 1,
        "K": 1024,
        "M": 1024**2,
        "G": 1024**3,
        "T": 1024**4,
        "P": 1024**5,
    }
    if suffix in multipliers:
        number = value[:-1] if suffix != "B" else value[:-1]
        if not number:
            number = "0"
        return int(float(number) * multipliers[suffix])
    try:
        return int(float(value))
    except ValueError:
        return None


def measure_disk_usage(
    weaviate_data_path: str,
    *,
    ssh_target: str | None = None,
    ssh_args: str | None = None,
) -> tuple[str, str, int | None]:
    target_path = str(weaviate_data_path or "").strip()
    if not target_path:
        raise RuntimeError("weaviate_data_path is required for disk usage measurement.")

    if ssh_target:
        ssh_bin = shutil.which("ssh")
        if not ssh_bin:
            raise RuntimeError(
                "ssh was not found on this machine. Set up OpenSSH or omit --disk-usage-ssh-target."
            )
        remote_cmd = f"du -sh {shlex.quote(target_path)}"
        extra_args = shlex.split(str(ssh_args or "").strip())
        cmd = [ssh_bin, *extra_args, ssh_target, remote_cmd]
        command_display = " ".join(
            [shlex.quote(ssh_bin), *[shlex.quote(arg) for arg in extra_args], shlex.quote(ssh_target), remote_cmd]
        )
    else:
        du_bin = shutil.which("du")
        if not du_bin:
            raise RuntimeError(
                "Local `du` was not found. When running the sweep from Windows against a remote cluster, "
                "pass --disk-usage-ssh-target <user@host-or-ssh-config-name>."
            )
        cmd = [du_bin, "-sh", target_path]
        command_display = f"du -sh {target_path}"

    result = subprocess.run(
        cmd,
        capture_output=True,
        text=True,
        check=True,
    )
    output = result.stdout.strip()
    if not output:
        raise RuntimeError(f"{command_display} returned no output for {target_path}")
    size_token = output.split()[0]
    return command_display, size_token, _human_size_to_bytes(size_token)


def upload_saved_embeddings(
    dataset_manifests: list[dict[str, Any]],
    *,
    batch_size: int,
    muvera_params: dict[str, int],
    pq_params: dict[str, Any],
    write_cluster_lmdb: bool,
) -> dict[str, Any]:
    total_uploaded = 0
    dataset_counts: list[dict[str, Any]] = []
    for manifest in dataset_manifests:
        storage_path = Path(manifest["storage_path"])
        storage_format = str(manifest.get("storage_format") or "jsonl.gz").strip().lower()
        dataset_id = str(manifest["canonical_dataset_id"])
        count = 0
        buffer: list[dict[str, Any]] = []
        if storage_format == "lmdb":
            records = load_dataset_embeddings_from_lmdb(storage_path)
        else:
            records = _load_saved_records_from_file(storage_path)
        for record in records:
            buffer.append(record)
            if len(buffer) >= batch_size:
                store_embeddings(
                    buffer,
                    muvera_params=muvera_params,
                    pq_params=pq_params,
                    write_lmdb=write_cluster_lmdb,
                )
                count += len(buffer)
                total_uploaded += len(buffer)
                buffer.clear()
        if buffer:
            store_embeddings(
                buffer,
                muvera_params=muvera_params,
                pq_params=pq_params,
                write_lmdb=write_cluster_lmdb,
            )
            count += len(buffer)
            total_uploaded += len(buffer)
            buffer.clear()
        dataset_counts.append(
            {
                "canonical_dataset_id": dataset_id,
                "uploaded_records": count,
                "storage_format": storage_format,
                "storage_path": str(storage_path),
            }
        )
        print(f"[phase2] Uploaded dataset {dataset_id}: {count} records")
    return {
        "total_uploaded_records": total_uploaded,
        "dataset_uploads": dataset_counts,
    }


def run_eval_suite(
    *,
    run_dir: Path,
    queries: Path,
    qrels: Path,
    eval_mode: str,
    retrieval_limit_list: str,
    k_list: str,
    shard: str,
    seed: int,
    cache_dir: Path | None,
    retrieval_mode: str,
    hybrid_alpha: float,
    rerank_source: str,
    rerank_k: int | None,
    diag_limit: int | None,
    diag_k: int,
    curve_kmax: int,
    heat_k: int,
    balanced_k: int,
    recall_threshold: float,
    latency_cap_ms: float | None,
    resume_existing_runs: bool = False,
) -> tuple[Path, list[dict[str, Any]], Path]:
    eval_output_dir = run_dir / "eval_suite"
    eval_output_dir.mkdir(parents=True, exist_ok=True)
    log_path = run_dir / "eval_suite.log"
    cmd = [
        sys.executable,
        "-u",
        str(Path(__file__).resolve().with_name("eval_suite.py")),
        "--queries",
        str(queries),
        "--qrels",
        str(qrels),
        "--mode",
        str(eval_mode),
        "--retrieval-limit-list",
        str(retrieval_limit_list),
        "--k-list",
        str(k_list),
        "--output-dir",
        str(eval_output_dir),
        "--shard",
        str(shard),
        "--seed",
        str(int(seed)),
        "--retrieval-mode",
        str(retrieval_mode),
        "--hybrid-alpha",
        str(float(hybrid_alpha)),
        "--rerank-source",
        str(rerank_source),
        "--diag-k",
        str(int(diag_k)),
        "--curve-kmax",
        str(int(curve_kmax)),
        "--heat-k",
        str(int(heat_k)),
        "--balanced-k",
        str(int(balanced_k)),
        "--recall-threshold",
        str(float(recall_threshold)),
    ]
    if rerank_k is not None:
        cmd.extend(["--rerank-k", str(int(rerank_k))])
    if cache_dir is not None:
        cmd.extend(["--cache-dir", str(cache_dir)])
    if diag_limit is not None:
        cmd.extend(["--diag-limit", str(int(diag_limit))])
    if latency_cap_ms is not None:
        cmd.extend(["--latency-cap-ms", str(float(latency_cap_ms))])
    if resume_existing_runs:
        cmd.append("--resume-existing-runs")

    log_mode = "a" if resume_existing_runs and log_path.exists() else "w"
    with log_path.open(log_mode, encoding="utf-8", buffering=1) as log_handle:
        if resume_existing_runs:
            banner = f"[phase2] Resuming eval from existing run files in {eval_output_dir}\n"
            print(banner, end="", flush=True)
            log_handle.write(banner)
        proc = subprocess.Popen(
            cmd,
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
            bufsize=1,
        )
        assert proc.stdout is not None
        for line in proc.stdout:
            print(line, end="", flush=True)
            log_handle.write(line)
            log_handle.flush()
        return_code = proc.wait()
        if return_code != 0:
            raise subprocess.CalledProcessError(return_code, cmd)
    summary_csv = eval_output_dir / "summary_metrics.csv"
    summary_rows = _read_csv_rows(summary_csv)
    return eval_output_dir, summary_rows, log_path


def _dataset_range(dataset_manifests: list[dict[str, Any]]) -> dict[str, Any]:
    dates = sorted(str(m.get("canonical_dataset_date") or "") for m in dataset_manifests)
    return {
        "dataset_ids": [str(m.get("canonical_dataset_id") or "") for m in dataset_manifests],
        "date_start": dates[0] if dates else "",
        "date_end": dates[-1] if dates else "",
    }


def _run_id_for_params(*, ksim: int, dprojections: int, repetitions: int, pq_enabled: bool) -> str:
    return f"pq{1 if pq_enabled else 0}_k{int(ksim)}_d{int(dprojections)}_r{int(repetitions)}"


def _load_completed_run_from_manifest(manifest_path: Path) -> dict[str, Any]:
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    summary_rows = manifest.get("evaluation", {}).get("summary_metrics") or []
    if not summary_rows:
        summary_csv = Path(
            manifest.get("evaluation", {}).get("summary_metrics_csv")
            or manifest_path.parent / "eval_suite" / "summary_metrics.csv"
        )
        summary_rows = _read_csv_rows(summary_csv)
    return {
        "manifest": manifest,
        "manifest_path": str(manifest_path.resolve()),
        "summary_rows": summary_rows,
    }


def run_single_configuration(
    *,
    output_root: Path,
    dataset_manifests: list[dict[str, Any]],
    queries: Path,
    qrels: Path,
    weaviate_data_path: str,
    disk_usage_ssh_target: str | None,
    disk_usage_ssh_args: str | None,
    skip_disk_usage: bool,
    upload_batch_size: int,
    eval_mode: str,
    retrieval_limit_list: str,
    k_list: str,
    shard: str,
    seed: int,
    cache_dir: Path | None,
    retrieval_mode: str,
    hybrid_alpha: float,
    rerank_source: str,
    rerank_k: int | None,
    diag_limit: int | None,
    diag_k: int,
    curve_kmax: int,
    heat_k: int,
    balanced_k: int,
    recall_threshold: float,
    latency_cap_ms: float | None,
    ksim: int,
    dprojections: int,
    repetitions: int,
    pq_enabled: bool,
    write_cluster_lmdb: bool,
) -> dict[str, Any]:
    run_id = _run_id_for_params(
        ksim=int(ksim),
        dprojections=int(dprojections),
        repetitions=int(repetitions),
        pq_enabled=bool(pq_enabled),
    )
    run_dir = output_root / "runs" / run_id
    run_dir.mkdir(parents=True, exist_ok=True)
    run_started_at = datetime.now().astimezone().isoformat()
    muvera_params = {
        "ksim": int(ksim),
        "dprojections": int(dprojections),
        "repetitions": int(repetitions),
    }
    pq_params = {
        "enabled": bool(pq_enabled),
        "centroids": int(WEAVIATE_PQ_CENTROIDS),
        "segments": int(WEAVIATE_PQ_SEGMENTS),
        "training_limit": int(WEAVIATE_PQ_TRAINING_LIMIT),
        "bit_compression": bool(WEAVIATE_PQ_BIT_COMPRESSION),
    }
    dimensionality = int(repetitions) * (2 ** int(ksim)) * int(dprojections)
    dataset_range = _dataset_range(dataset_manifests)

    total_start = time.perf_counter()
    upload_checkpoint_path = _upload_checkpoint_path(run_dir)
    upload_checkpoint = _json_load(upload_checkpoint_path) if upload_checkpoint_path.exists() else None
    existing_eval_runs = _existing_eval_run_files(run_dir)
    resume_eval_only = bool(upload_checkpoint or existing_eval_runs)
    resume_reason = ""
    resume_validation_note = ""
    stale_resume_artifacts = False

    if resume_eval_only:
        live_claim_count = _current_claim_total_count()
        expected_uploaded_records = None
        if upload_checkpoint is not None:
            expected_uploaded_records = int(
                ((upload_checkpoint.get("upload") or {}).get("total_uploaded_records") or 0)
            )

        if live_claim_count == 0:
            resume_validation_note = (
                "Found prior upload/eval artifacts, but the live Claim collection is empty. "
                "Falling back to full reset + upload."
            )
            print(f"[phase2] {resume_validation_note}")
            resume_eval_only = False
            stale_resume_artifacts = True
        elif (
            live_claim_count is not None
            and expected_uploaded_records is not None
            and expected_uploaded_records > 0
            and live_claim_count != expected_uploaded_records
        ):
            resume_validation_note = (
                "Found prior upload checkpoint, but live Claim object count "
                f"({live_claim_count}) does not match expected uploaded record count "
                f"({expected_uploaded_records}). Falling back to full reset + upload."
            )
            print(f"[phase2] {resume_validation_note}")
            resume_eval_only = False
            stale_resume_artifacts = True

    if stale_resume_artifacts:
        upload_checkpoint = None
        existing_eval_runs = []
        if upload_checkpoint_path.exists():
            try:
                upload_checkpoint_path.unlink()
                print(f"[phase2] Removed stale upload checkpoint: {upload_checkpoint_path}")
            except Exception as exc:
                print(f"[warn] Failed to remove stale upload checkpoint {upload_checkpoint_path}: {exc}")
        stale_eval_dir = run_dir / "eval_suite"
        stale_eval_log = run_dir / "eval_suite.log"
        if stale_eval_dir.exists():
            try:
                shutil.rmtree(stale_eval_dir, ignore_errors=True)
                print(f"[phase2] Removed stale eval directory: {stale_eval_dir}")
            except Exception as exc:
                print(f"[warn] Failed to remove stale eval directory {stale_eval_dir}: {exc}")
        if stale_eval_log.exists():
            try:
                stale_eval_log.unlink()
                print(f"[phase2] Removed stale eval log: {stale_eval_log}")
            except Exception as exc:
                print(f"[warn] Failed to remove stale eval log {stale_eval_log}: {exc}")

    if resume_eval_only:
        if upload_checkpoint is not None:
            resume_reason = "upload_checkpoint"
            deleted_collections = list(upload_checkpoint.get("deleted_collections") or [])
            reset_seconds = upload_checkpoint.get("reset_seconds")
            upload_seconds = upload_checkpoint.get("upload_seconds")
            upload_stats = upload_checkpoint.get("upload") or _build_upload_stats_from_dataset_manifests(dataset_manifests)
            checkpoint_disk = upload_checkpoint.get("disk_usage") or {}
            disk_usage_command = str(checkpoint_disk.get("command") or "")
            disk_usage_raw = str(checkpoint_disk.get("raw") or "")
            disk_usage_bytes = checkpoint_disk.get("bytes")
            disk_usage_status = str(checkpoint_disk.get("status") or "")
            disk_usage_note = str(checkpoint_disk.get("note") or "")
        else:
            resume_reason = "existing_eval_run_files"
            deleted_collections = []
            reset_seconds = None
            upload_seconds = None
            upload_stats = _build_upload_stats_from_dataset_manifests(dataset_manifests)
            disk_usage_command = ""
            disk_usage_raw = ""
            disk_usage_bytes = None
            disk_usage_status = ""
            disk_usage_note = (
                "Upload checkpoint was unavailable, so upload/reset timings could not be recovered. "
                "Existing eval run files were used to resume from the current Weaviate state."
            )

        if skip_disk_usage:
            if not disk_usage_status:
                disk_usage_status = "skipped"
                disk_usage_note = (
                    "Disk usage measurement was skipped because the sweep controller could not execute "
                    "`du -sh` on the Weaviate host."
                )
        elif not disk_usage_status:
            disk_usage_command, disk_usage_raw, disk_usage_bytes = measure_disk_usage(
                weaviate_data_path,
                ssh_target=disk_usage_ssh_target,
                ssh_args=disk_usage_ssh_args,
            )
            disk_usage_status = "measured"
            disk_usage_note = ""

        print(
            f"[phase2] Resuming eval-only for run {run_id}: reason={resume_reason} "
            f"existing_run_files={len(existing_eval_runs)}"
        )
    else:
        reset_start = time.perf_counter()
        deleted_collections = reset_weaviate_state(muvera_params=muvera_params, pq_params=pq_params)
        reset_seconds = time.perf_counter() - reset_start

        upload_start = time.perf_counter()
        upload_stats = upload_saved_embeddings(
            dataset_manifests,
            batch_size=int(upload_batch_size),
            muvera_params=muvera_params,
            pq_params=pq_params,
            write_cluster_lmdb=write_cluster_lmdb,
        )
        upload_seconds = time.perf_counter() - upload_start

        if skip_disk_usage:
            disk_usage_command = ""
            disk_usage_raw = ""
            disk_usage_bytes = None
            disk_usage_status = "skipped"
            disk_usage_note = (
                "Disk usage measurement was skipped because the sweep controller could not execute "
                "`du -sh` on the Weaviate host."
            )
        else:
            disk_usage_command, disk_usage_raw, disk_usage_bytes = measure_disk_usage(
                weaviate_data_path,
                ssh_target=disk_usage_ssh_target,
                ssh_args=disk_usage_ssh_args,
            )
            disk_usage_status = "measured"
            disk_usage_note = ""

        _json_dump(
            upload_checkpoint_path,
            {
                "schema_version": 1,
                "phase": "parameter_sweep_upload_complete",
                "run_id": run_id,
                "timestamp": run_started_at,
                "deleted_collections": list(deleted_collections),
                "reset_seconds": reset_seconds,
                "upload_seconds": upload_seconds,
                "upload": upload_stats,
                "disk_usage": {
                    "status": disk_usage_status,
                    "command": disk_usage_command,
                    "raw": disk_usage_raw,
                    "bytes": disk_usage_bytes,
                    "path": str(weaviate_data_path),
                    "note": disk_usage_note,
                },
            },
        )

        client_claim_count = _current_claim_total_count()
        graphql_claim_count, graphql_endpoint = _graphql_claim_total_count()
        expected_uploaded_records = int((upload_stats or {}).get("total_uploaded_records") or 0)
        print(
            f"[phase2] Post-upload counts: client_claim_count={client_claim_count} "
            f"graphql_claim_count={graphql_claim_count} graphql={graphql_endpoint}"
        )
        if expected_uploaded_records > 0:
            if client_claim_count == 0:
                raise RuntimeError(
                    "Upload completed without surfacing an insert error, but the live Weaviate client sees zero Claim objects. "
                    "This points to a failed or misdirected upload."
                )
            if graphql_claim_count == 0:
                raise RuntimeError(
                    "Upload completed, but the GraphQL endpoint used by eval sees zero Claim objects. "
                    "Upload and eval are likely targeting different Weaviate endpoints. "
                    f"GraphQL endpoint: {graphql_endpoint}"
                )
            if (
                client_claim_count is not None
                and graphql_claim_count is not None
                and client_claim_count != graphql_claim_count
            ):
                raise RuntimeError(
                    "Upload completed, but Weaviate client count and GraphQL count do not match. "
                    "Upload and eval may be targeting different endpoints. "
                    f"client_claim_count={client_claim_count} graphql_claim_count={graphql_claim_count} "
                    f"graphql={graphql_endpoint}"
                )

    print(f"[phase2] Starting eval for run {run_id}")
    eval_start = time.perf_counter()
    eval_output_dir, summary_rows, eval_log_path = run_eval_suite(
        run_dir=run_dir,
        queries=queries,
        qrels=qrels,
        eval_mode=eval_mode,
        retrieval_limit_list=retrieval_limit_list,
        k_list=k_list,
        shard=shard,
        seed=seed,
        cache_dir=cache_dir,
        retrieval_mode=retrieval_mode,
        hybrid_alpha=hybrid_alpha,
        rerank_source=rerank_source,
        rerank_k=rerank_k,
        diag_limit=diag_limit,
        diag_k=diag_k,
        curve_kmax=curve_kmax,
        heat_k=heat_k,
        balanced_k=balanced_k,
        recall_threshold=recall_threshold,
        latency_cap_ms=latency_cap_ms,
        resume_existing_runs=resume_eval_only,
    )
    eval_seconds = time.perf_counter() - eval_start
    if resume_eval_only:
        total_seconds = float(reset_seconds or 0.0) + float(upload_seconds or 0.0) + float(eval_seconds)
    else:
        total_seconds = time.perf_counter() - total_start

    manifest = {
        "schema_version": 1,
        "phase": "parameter_sweep",
        "run_id": run_id,
        "timestamp": run_started_at,
        "dataset": dataset_range,
        "parameters": {
            **muvera_params,
            "pq_enabled": bool(pq_enabled),
        },
        "pq": pq_params,
        "computed_dimensionality": dimensionality,
        "disk_usage": {
            "status": disk_usage_status,
            "command": disk_usage_command,
            "raw": disk_usage_raw,
            "bytes": disk_usage_bytes,
            "path": str(weaviate_data_path),
            "note": disk_usage_note,
        },
        "runtime": {
            "reset_seconds": reset_seconds,
            "upload_seconds": upload_seconds,
            "evaluation_seconds": eval_seconds,
            "total_seconds": total_seconds,
        },
        "resume": {
            "eval_only_resumed": bool(resume_eval_only),
            "reason": resume_reason,
            "upload_checkpoint_path": str(upload_checkpoint_path.resolve()) if upload_checkpoint_path.exists() else "",
            "existing_eval_run_files": [p.name for p in existing_eval_runs],
            "note": (
                (resume_validation_note + " " if resume_validation_note else "")
                + "Eval resumed from existing Weaviate state and prior .run files."
                if resume_eval_only
                else resume_validation_note
            ),
        },
        "reset_behavior": {
            "collections_deleted": list(deleted_collections),
            "configured_default": list(PARAMETER_SWEEP_RESET_COLLECTIONS),
            "note": (
                "Patent and Claim are both reset per run because store_embeddings repopulates both "
                "collections and disk usage is measured on the full Weaviate data directory."
            ),
        },
        "upload": upload_stats,
        "cluster_lmdb": {
            "enabled": bool(write_cluster_lmdb),
            "note": (
                "Cluster-side LMDB side stores are scratch only. "
                "Leave disabled for sweep runs unless explicitly needed."
            ),
        },
        "evaluation": {
            "queries": str(queries),
            "qrels": str(qrels),
            "mode": eval_mode,
            "retrieval_limit_list": retrieval_limit_list,
            "k_list": k_list,
            "shard": shard,
            "seed": seed,
            "retrieval_mode": retrieval_mode,
            "hybrid_alpha": hybrid_alpha,
            "rerank_source": rerank_source,
            "rerank_k": rerank_k,
            "diag_limit": diag_limit,
            "diag_k": diag_k,
            "curve_kmax": curve_kmax,
            "heat_k": heat_k,
            "balanced_k": balanced_k,
            "recall_threshold": recall_threshold,
            "latency_cap_ms": latency_cap_ms,
            "output_dir": str(eval_output_dir),
            "summary_metrics_csv": str((eval_output_dir / "summary_metrics.csv").resolve()),
            "per_query_metrics_csv": str((eval_output_dir / "per_query_metrics.csv").resolve()),
            "log_path": str(eval_log_path.resolve()),
            "summary_metrics": summary_rows,
        },
    }
    manifest_path = run_dir / "manifest.json"
    _json_dump(manifest_path, manifest)
    print(
        f"[phase2] Completed run {run_id}: pq={int(bool(pq_enabled))} disk={disk_usage_raw} dim={dimensionality} "
        f"upload_s={_format_seconds(upload_seconds)} eval_s={_format_seconds(eval_seconds)}"
    )
    return {
        "manifest": manifest,
        "manifest_path": str(manifest_path.resolve()),
        "summary_rows": summary_rows,
    }


def export_parameter_sweep_excel(
    runs: list[dict[str, Any]],
    *,
    workbook_path: Path,
) -> None:
    wb = Workbook()
    master = wb.active
    master.title = "master"

    overview_headers = [
        "run_id",
        "timestamp",
        "dataset_ids",
        "date_start",
        "date_end",
        "pq_enabled",
        "pq_centroids",
        "pq_segments",
        "pq_training_limit",
        "pq_bit_compression",
        "ksim",
        "dprojections",
        "repetitions",
        "computed_dimensionality",
        "disk_usage_status",
        "disk_usage_raw",
        "disk_usage_bytes",
        "reset_seconds",
        "upload_seconds",
        "evaluation_seconds",
        "total_seconds",
        "retrieval_mode",
        "hybrid_alpha",
        "rerank_source",
        "cluster_lmdb_enabled",
        "manifest_path",
    ]
    overview_rows: list[dict[str, Any]] = []
    summary_rows: list[dict[str, Any]] = []

    for run in runs:
        manifest = run["manifest"]
        params = manifest["parameters"]
        dataset = manifest["dataset"]
        runtime = manifest["runtime"]
        disk = manifest["disk_usage"]
        evaluation = manifest["evaluation"]
        overview_rows.append(
            {
                "run_id": manifest["run_id"],
                "timestamp": manifest["timestamp"],
                "dataset_ids": ", ".join(dataset["dataset_ids"]),
                "date_start": dataset["date_start"],
                "date_end": dataset["date_end"],
                "pq_enabled": manifest["parameters"]["pq_enabled"],
                "pq_centroids": manifest["pq"]["centroids"],
                "pq_segments": manifest["pq"]["segments"],
                "pq_training_limit": manifest["pq"]["training_limit"],
                "pq_bit_compression": manifest["pq"]["bit_compression"],
                "ksim": params["ksim"],
                "dprojections": params["dprojections"],
                "repetitions": params["repetitions"],
                "computed_dimensionality": manifest["computed_dimensionality"],
                "disk_usage_status": disk["status"],
                "disk_usage_raw": disk["raw"],
                "disk_usage_bytes": disk["bytes"],
                "reset_seconds": runtime["reset_seconds"],
                "upload_seconds": runtime["upload_seconds"],
                "evaluation_seconds": runtime["evaluation_seconds"],
                "total_seconds": runtime["total_seconds"],
                "retrieval_mode": evaluation["retrieval_mode"],
                "hybrid_alpha": evaluation["hybrid_alpha"],
                "rerank_source": evaluation["rerank_source"],
                "cluster_lmdb_enabled": manifest["cluster_lmdb"]["enabled"],
                "manifest_path": run["manifest_path"],
            }
        )
        for row in run["summary_rows"]:
            summary_rows.append(
                {
                    "run_id": manifest["run_id"],
                    "pq_enabled": manifest["parameters"]["pq_enabled"],
                    "ksim": params["ksim"],
                    "dprojections": params["dprojections"],
                    "repetitions": params["repetitions"],
                    "computed_dimensionality": manifest["computed_dimensionality"],
                    "disk_usage_status": disk["status"],
                    "disk_usage_raw": disk["raw"],
                    "disk_usage_bytes": disk["bytes"],
                    **row,
                }
            )

    cursor = _write_table(master, 1, overview_headers, overview_rows)
    cursor += 2
    if summary_rows:
        summary_headers = list(summary_rows[0].keys())
        cursor = _write_table(master, cursor, summary_headers, summary_rows)
        _auto_fit_columns(master, max_col=len(summary_headers), max_row=cursor)
    else:
        _auto_fit_columns(master, max_col=len(overview_headers), max_row=cursor)

    master.freeze_panes = "A2"

    for run in runs:
        manifest = run["manifest"]
        sheet_name = manifest["run_id"][:31]
        ws = wb.create_sheet(title=sheet_name)

        metadata_rows = [
            {"field": "run_id", "value": manifest["run_id"]},
            {"field": "timestamp", "value": manifest["timestamp"]},
            {"field": "dataset_ids", "value": ", ".join(manifest["dataset"]["dataset_ids"])},
            {"field": "date_start", "value": manifest["dataset"]["date_start"]},
            {"field": "date_end", "value": manifest["dataset"]["date_end"]},
            {"field": "pq_enabled", "value": manifest["parameters"]["pq_enabled"]},
            {"field": "pq_centroids", "value": manifest["pq"]["centroids"]},
            {"field": "pq_segments", "value": manifest["pq"]["segments"]},
            {"field": "pq_training_limit", "value": manifest["pq"]["training_limit"]},
            {"field": "pq_bit_compression", "value": manifest["pq"]["bit_compression"]},
            {"field": "ksim", "value": manifest["parameters"]["ksim"]},
            {"field": "dprojections", "value": manifest["parameters"]["dprojections"]},
            {"field": "repetitions", "value": manifest["parameters"]["repetitions"]},
            {"field": "computed_dimensionality", "value": manifest["computed_dimensionality"]},
            {"field": "disk_usage_status", "value": manifest["disk_usage"]["status"]},
            {"field": "disk_usage_raw", "value": manifest["disk_usage"]["raw"]},
            {"field": "disk_usage_bytes", "value": manifest["disk_usage"]["bytes"]},
            {"field": "disk_usage_note", "value": manifest["disk_usage"]["note"]},
            {"field": "reset_seconds", "value": manifest["runtime"]["reset_seconds"]},
            {"field": "upload_seconds", "value": manifest["runtime"]["upload_seconds"]},
            {"field": "evaluation_seconds", "value": manifest["runtime"]["evaluation_seconds"]},
            {"field": "total_seconds", "value": manifest["runtime"]["total_seconds"]},
            {"field": "retrieval_mode", "value": manifest["evaluation"]["retrieval_mode"]},
            {"field": "hybrid_alpha", "value": manifest["evaluation"]["hybrid_alpha"]},
            {"field": "rerank_source", "value": manifest["evaluation"]["rerank_source"]},
            {"field": "cluster_lmdb_enabled", "value": manifest["cluster_lmdb"]["enabled"]},
            {"field": "manifest_path", "value": run["manifest_path"]},
            {"field": "summary_metrics_csv", "value": manifest["evaluation"]["summary_metrics_csv"]},
            {"field": "per_query_metrics_csv", "value": manifest["evaluation"]["per_query_metrics_csv"]},
        ]
        cursor = _write_table(ws, 1, ["field", "value"], metadata_rows)
        cursor += 2

        dataset_rows = []
        for item in manifest["upload"]["dataset_uploads"]:
            dataset_rows.append(
                {
                    "canonical_dataset_id": item["canonical_dataset_id"],
                    "uploaded_records": item["uploaded_records"],
                    "storage_format": item["storage_format"],
                    "storage_path": item["storage_path"],
                }
            )
        if dataset_rows:
            cursor = _write_table(ws, cursor, list(dataset_rows[0].keys()), dataset_rows)
            cursor += 2

        if run["summary_rows"]:
            cursor = _write_table(ws, cursor, list(run["summary_rows"][0].keys()), run["summary_rows"])

        ws.freeze_panes = "A2"
        _auto_fit_columns(ws, max_col=ws.max_column, max_row=ws.max_row)

    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(workbook_path)
    print(f"[phase2] Wrote workbook: {workbook_path}")


def _build_sweep_manifest(
    *,
    args: argparse.Namespace,
    output_root: Path,
    dataset_ids: list[str],
    runs: list[dict[str, Any]],
    ksim_values: list[int],
    dprojection_values: list[int],
    repetition_values: list[int],
    pq_enabled_values: list[bool],
    skip_disk_usage: bool,
) -> dict[str, Any]:
    return {
        "schema_version": 1,
        "created_at": datetime.now().astimezone().isoformat(),
        "phase": "parameter_sweep",
        "embeddings_roots": [str(Path(root).resolve()) for root in args.embeddings_root],
        "output_root": str(output_root),
        "dataset_ids": dataset_ids,
        "weaviate_data_path": str(args.weaviate_data_path),
        "disk_usage_ssh_target": str(args.disk_usage_ssh_target or ""),
        "disk_usage_ssh_args": str(args.disk_usage_ssh_args or ""),
        "skip_disk_usage": skip_disk_usage,
        "queries": str(Path(args.queries).resolve()),
        "qrels": str(Path(args.qrels).resolve()),
        "write_cluster_lmdb": bool(args.write_cluster_lmdb),
        "pq_enabled_values": pq_enabled_values,
        "pq_defaults": {
            "enabled": bool(WEAVIATE_PQ_ENABLED),
            "centroids": int(WEAVIATE_PQ_CENTROIDS),
            "segments": int(WEAVIATE_PQ_SEGMENTS),
            "training_limit": int(WEAVIATE_PQ_TRAINING_LIMIT),
            "bit_compression": bool(WEAVIATE_PQ_BIT_COMPRESSION),
        },
        "ksim_values": ksim_values,
        "dprojection_values": dprojection_values,
        "repetition_values": repetition_values,
        "run_ids": [run["manifest"]["run_id"] for run in runs],
    }


def _write_sweep_outputs(
    *,
    args: argparse.Namespace,
    output_root: Path,
    dataset_ids: list[str],
    runs: list[dict[str, Any]],
    ksim_values: list[int],
    dprojection_values: list[int],
    repetition_values: list[int],
    pq_enabled_values: list[bool],
    skip_disk_usage: bool,
) -> None:
    sweep_manifest = _build_sweep_manifest(
        args=args,
        output_root=output_root,
        dataset_ids=dataset_ids,
        runs=runs,
        ksim_values=ksim_values,
        dprojection_values=dprojection_values,
        repetition_values=repetition_values,
        pq_enabled_values=pq_enabled_values,
        skip_disk_usage=skip_disk_usage,
    )
    _json_dump(output_root / "sweep_manifest.json", sweep_manifest)
    export_parameter_sweep_excel(runs, workbook_path=output_root / "parameter sweep.xlsx")


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Run MUVERA parameter sweeps from saved embedding datasets without re-embedding."
    )
    parser.add_argument(
        "--embeddings-root",
        type=Path,
        nargs="+",
        required=True,
        help="One or more roots containing per-dataset manifest.json + embedding stores.",
    )
    parser.add_argument("--output-root", type=Path, required=True)
    parser.add_argument("--start-date", type=str, default="2025-09-02")
    parser.add_argument("--num-datasets", type=int, default=4)
    parser.add_argument("--dataset-ids", type=str, default="")
    parser.add_argument("--download-path", type=Path, default=DOWNLOAD_DIR)
    parser.add_argument("--dataset-product", type=str, default="PTGRDT")
    parser.add_argument("--ksim-list", type=str, required=True)
    parser.add_argument("--dprojections-list", type=str, required=True)
    parser.add_argument("--repetitions-list", type=str, required=True)
    parser.add_argument(
        "--pq-enabled-list",
        type=str,
        default="true" if WEAVIATE_PQ_ENABLED else "false",
        help="Comma-separated PQ enabled states, e.g. true,false",
    )
    parser.add_argument("--upload-batch-size", type=int, default=256)
    parser.add_argument("--weaviate-data-path", type=str, required=True)
    parser.add_argument(
        "--disk-usage-ssh-target",
        type=str,
        default=WEAVIATE_DISK_USAGE_SSH_TARGET,
        help=(
            "Optional SSH target used to run `du -sh <weaviate_data_path>` on the cluster. "
            "Required when the sweep controller is not running on the same Linux host as Weaviate."
        ),
    )
    parser.add_argument(
        "--disk-usage-ssh-args",
        type=str,
        default=WEAVIATE_DISK_USAGE_SSH_ARGS,
        help=(
            "Optional extra ssh args for disk usage measurement, for example "
            "\"-J user@login-host\" or \"-i C:/path/to/key\"."
        ),
    )
    parser.add_argument(
        "--skip-disk-usage",
        action="store_true",
        help="Skip `du -sh` measurement and record disk usage as unavailable in run manifests and Excel.",
    )
    parser.add_argument("--queries", type=Path, required=True)
    parser.add_argument("--qrels", type=Path, required=True)
    parser.add_argument("--eval-mode", type=str, default="both", choices=["retrieve", "rerank", "both"])
    parser.add_argument("--retrieval-limit-list", type=str, default="100,200,400,800,1200,2000")
    parser.add_argument("--k-list", type=str, default="10,20,50,100,200")
    parser.add_argument("--rerank-k", type=int, default=None)
    parser.add_argument("--shard", type=str, default="128_f16")
    parser.add_argument("--seed", type=int, default=13)
    parser.add_argument("--cache-dir", type=Path, default=None)
    parser.add_argument("--retrieval-mode", type=str, default="vector", choices=["vector", "bm25", "hybrid"])
    parser.add_argument("--hybrid-alpha", type=float, default=0.5)
    parser.add_argument("--rerank-source", type=str, default="weaviate", choices=["lmdb", "weaviate"])
    parser.add_argument(
        "--write-cluster-lmdb",
        action="store_true",
        help="Also write local LMDB side stores during upload. Default is off for scratch-only sweeps.",
    )
    parser.add_argument("--diag-limit", type=int, default=None)
    parser.add_argument("--diag-k", type=int, default=200)
    parser.add_argument("--curve-kmax", type=int, default=200)
    parser.add_argument("--heat-k", type=int, default=200)
    parser.add_argument("--balanced-k", type=int, default=50)
    parser.add_argument("--recall-threshold", type=float, default=0.70)
    parser.add_argument("--latency-cap-ms", type=float, default=None)
    parser.add_argument(
        "--rerun-completed",
        action="store_true",
        help="Ignore existing per-run manifests and rerun completed configurations from scratch.",
    )
    args = parser.parse_args()

    dataset_ids = _resolve_dataset_ids(
        start_date=args.start_date,
        num_datasets=int(args.num_datasets),
        dataset_ids_raw=args.dataset_ids,
        download_path=Path(args.download_path).resolve(),
        dataset_product=args.dataset_product,
    )
    embeddings_roots = [Path(root).resolve() for root in args.embeddings_root]
    dataset_manifests = _load_dataset_manifests(embeddings_roots, dataset_ids)
    ksim_values = _parse_int_list(args.ksim_list, "--ksim-list")
    dprojection_values = _parse_int_list(args.dprojections_list, "--dprojections-list")
    repetition_values = _parse_int_list(args.repetitions_list, "--repetitions-list")
    pq_enabled_values = _parse_bool_list(args.pq_enabled_list, "--pq-enabled-list")
    if args.rerank_source == "lmdb" and not args.write_cluster_lmdb:
        raise SystemExit(
            "--rerank-source lmdb requires --write-cluster-lmdb so the cluster has local rerank side stores."
        )
    disk_usage_ssh_target = str(args.disk_usage_ssh_target).strip() or None
    skip_disk_usage = bool(args.skip_disk_usage)
    if skip_disk_usage:
        print("[phase2] Disk usage measurement disabled by --skip-disk-usage")
    else:
        try:
            probe_command, probe_raw, _ = measure_disk_usage(
                str(args.weaviate_data_path),
                ssh_target=disk_usage_ssh_target,
                ssh_args=(str(args.disk_usage_ssh_args).strip() or None),
            )
        except Exception as exc:
            raise SystemExit(
                f"Disk usage probe failed before sweep start: {exc}. "
                "If you cannot run `du -sh` on the cluster host, rerun with --skip-disk-usage."
            ) from exc
        print(f"[phase2] Disk usage probe OK: {probe_command} -> {probe_raw}")

    output_root = Path(args.output_root).resolve()
    output_root.mkdir(parents=True, exist_ok=True)

    runs: list[dict[str, Any]] = []
    for ksim, dprojections, repetitions, pq_enabled in product(
        ksim_values,
        dprojection_values,
        repetition_values,
        pq_enabled_values,
    ):
        run_id = _run_id_for_params(
            ksim=int(ksim),
            dprojections=int(dprojections),
            repetitions=int(repetitions),
            pq_enabled=bool(pq_enabled),
        )
        manifest_path = output_root / "runs" / run_id / "manifest.json"
        if manifest_path.exists() and not bool(args.rerun_completed):
            print(f"[phase2] Reusing completed run: {run_id}")
            runs.append(_load_completed_run_from_manifest(manifest_path))
            continue
        if manifest_path.parent.exists() and not manifest_path.exists():
            print(f"[phase2] Found incomplete prior run; attempting resume: {run_id}")
        run = run_single_configuration(
            output_root=output_root,
            dataset_manifests=dataset_manifests,
            queries=Path(args.queries).resolve(),
            qrels=Path(args.qrels).resolve(),
            weaviate_data_path=str(args.weaviate_data_path),
            disk_usage_ssh_target=disk_usage_ssh_target,
            disk_usage_ssh_args=(str(args.disk_usage_ssh_args).strip() or None),
            skip_disk_usage=skip_disk_usage,
            upload_batch_size=int(args.upload_batch_size),
            eval_mode=args.eval_mode,
            retrieval_limit_list=args.retrieval_limit_list,
            k_list=args.k_list,
            shard=args.shard,
            seed=int(args.seed),
            cache_dir=Path(args.cache_dir).resolve() if args.cache_dir is not None else None,
            retrieval_mode=args.retrieval_mode,
            hybrid_alpha=float(args.hybrid_alpha),
            rerank_source=args.rerank_source,
            rerank_k=args.rerank_k,
            diag_limit=args.diag_limit,
            diag_k=int(args.diag_k),
            curve_kmax=int(args.curve_kmax),
            heat_k=int(args.heat_k),
            balanced_k=int(args.balanced_k),
            recall_threshold=float(args.recall_threshold),
            latency_cap_ms=args.latency_cap_ms,
            ksim=int(ksim),
            dprojections=int(dprojections),
            repetitions=int(repetitions),
            pq_enabled=bool(pq_enabled),
            write_cluster_lmdb=bool(args.write_cluster_lmdb),
        )
        runs.append(run)
        _write_sweep_outputs(
            args=args,
            output_root=output_root,
            dataset_ids=dataset_ids,
            runs=runs,
            ksim_values=ksim_values,
            dprojection_values=dprojection_values,
            repetition_values=repetition_values,
            pq_enabled_values=pq_enabled_values,
            skip_disk_usage=skip_disk_usage,
        )

    _write_sweep_outputs(
        args=args,
        output_root=output_root,
        dataset_ids=dataset_ids,
        runs=runs,
        ksim_values=ksim_values,
        dprojection_values=dprojection_values,
        repetition_values=repetition_values,
        pq_enabled_values=pq_enabled_values,
        skip_disk_usage=skip_disk_usage,
    )


if __name__ == "__main__":
    main()
