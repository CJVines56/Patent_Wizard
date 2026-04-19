"""
Run qrels evaluation across multiple retrieval modes and export an Excel workbook:
- one worksheet per run configuration
- one master worksheet with aggregate metrics + charts

Configurations executed:
- BM25 only
- Vector only
- Hybrid alpha 0.2
- Hybrid alpha 0.5
- Hybrid alpha 0.8
"""

from __future__ import annotations

import argparse
import csv
import os
import tempfile
from pathlib import Path
from typing import Any

try:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Font, PatternFill
except ImportError as e:
    raise SystemExit(
        "openpyxl is required for Excel export. Install with: poetry run python -m pip install openpyxl"
    ) from e

try:
    import matplotlib.pyplot as plt
except ImportError as e:
    raise SystemExit(
        "matplotlib is required for plotting. Install with: poetry run python -m pip install matplotlib"
    ) from e


RUNS = [
    {"sheet": "bm25", "retrieval_mode": "bm25", "hybrid_alpha": 0.0},
    {"sheet": "vector", "retrieval_mode": "vector", "hybrid_alpha": 1.0},
    {"sheet": "hybrid_a02", "retrieval_mode": "hybrid", "hybrid_alpha": 0.2},
    {"sheet": "hybrid_a05", "retrieval_mode": "hybrid", "hybrid_alpha": 0.5},
    {"sheet": "hybrid_a08", "retrieval_mode": "hybrid", "hybrid_alpha": 0.8},
]


def _read_csv_rows(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8", newline="") as f:
        reader = csv.DictReader(f)
        return [dict(row) for row in reader]


def _write_table(ws, start_row: int, headers: list[str], rows: list[dict[str, Any]]) -> int:
    header_fill = PatternFill(fill_type="solid", fgColor="D9E1F2")
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col_idx, value=header)
        cell.font = Font(bold=True)
        cell.fill = header_fill

    row_cursor = start_row + 1
    for row in rows:
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=row_cursor, column=col_idx, value=row.get(header, ""))
        row_cursor += 1
    return row_cursor


def _auto_fit_columns(ws, max_col: int, max_row: int) -> None:
    for col_idx in range(1, max_col + 1):
        col_letter = ws.cell(row=1, column=col_idx).column_letter
        width = 12
        for row_idx in range(1, max_row + 1):
            value = ws.cell(row=row_idx, column=col_idx).value
            if value is None:
                continue
            width = max(width, min(80, len(str(value)) + 2))
        ws.column_dimensions[col_letter].width = width


def _resolve_style_path(style_path: Path | None) -> Path | None:
    if style_path is not None:
        p = Path(style_path)
        if p.exists():
            return p
        raise FileNotFoundError(f"Matplotlib style file not found: {p}")

    env_raw = os.environ.get("MPLSTYLE_PATH", "").strip()
    style_env = Path(env_raw) if env_raw else None
    if style_env and style_env.exists():
        return style_env

    candidates = [
        Path("backend/validation/plot_style.mplstyle"),
        Path("backend/validation/matplotlib.mplstyle"),
        Path("backend/validation/style.mplstyle"),
        Path("backend/app/plot_style.mplstyle"),
        Path("backend/app/matplotlib.mplstyle"),
        Path("plot_style.mplstyle"),
        Path("matplotlib.mplstyle"),
        Path("style.mplstyle"),
    ]
    for c in candidates:
        if c.exists():
            return c
    return None


def _build_master_matplotlib_plot(
    run_outputs: list[dict[str, Any]],
    metric_headers: list[str],
    out_png: Path,
    style_path: Path | None,
) -> None:
    if style_path is not None:
        plt.style.use(str(style_path))

    x_labels = [str(r["run"]) for r in run_outputs]
    x = list(range(len(x_labels)))

    fig, ax = plt.subplots(figsize=(11, 6))
    for metric in metric_headers:
        y = [float(r.get(metric, 0.0)) for r in run_outputs]
        ax.plot(x, y, marker="o", linewidth=2, label=metric)

    ax.set_title("Qrels Metrics by Retrieval Setting")
    ax.set_xlabel("Run")
    ax.set_ylabel("Score")
    ax.set_xticks(x)
    ax.set_xticklabels(x_labels, rotation=15, ha="right")
    ax.set_ylim(0.0, 1.0)
    ax.grid(True, linestyle="--", alpha=0.35)
    ax.legend(loc="best")
    fig.tight_layout()
    out_png.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(out_png, dpi=160)
    plt.close(fig)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Run BM25/vector/hybrid mode sweep on qrels and export an Excel workbook."
    )
    parser.add_argument("--queries", type=Path, required=True, help="Path to queries JSONL.")
    parser.add_argument("--qrels", type=Path, required=True, help="Path to qrels JSONL.")
    parser.add_argument("--output", type=Path, required=True, help="Output XLSX path.")
    parser.add_argument("--retrieve-shard", type=str, default="128_f16")
    parser.add_argument("--rerank-shard", type=str, default="128_f16")
    parser.add_argument("--limit", type=int, default=400)
    parser.add_argument("--rerank-k", type=int, default=200)
    parser.add_argument("--rerank-source", type=str, default="lmdb", choices=["lmdb", "weaviate"])
    parser.add_argument(
        "--mplstyle",
        type=Path,
        default=None,
        help="Optional matplotlib style file (.mplstyle). If omitted, common project paths are checked.",
    )
    parser.add_argument(
        "--filter-missing-qrels",
        action=argparse.BooleanOptionalAction,
        default=False,
        help="Drop queries whose relevant claims are not in index (default: off).",
    )
    args = parser.parse_args()

    from backend.app.scripts.retrieve_rerank import evaluate

    if not args.queries.exists():
        raise SystemExit(f"Missing queries file: {args.queries}")
    if not args.qrels.exists():
        raise SystemExit(f"Missing qrels file: {args.qrels}")

    candidate_metric = f"candidate_hit@{int(args.limit)}"
    metric_headers = ["precision@10", "recall@10", "ndcg@10", "mrr@10", candidate_metric]

    run_outputs: list[dict[str, Any]] = []
    per_query_by_run: dict[str, list[dict[str, Any]]] = {}
    style_path = _resolve_style_path(args.mplstyle)
    if style_path is not None:
        print(f"[plot] using matplotlib style: {style_path}")
    else:
        print("[plot] no matplotlib style file found; using default style")
    plot_png_path = args.output.parent / f"{args.output.stem}_master_plot.png"

    with tempfile.TemporaryDirectory(prefix="qrels_sweep_") as td:
        tmpdir = Path(td)
        for cfg in RUNS:
            sheet = cfg["sheet"]
            per_query_csv = tmpdir / f"{sheet}_per_query.csv"
            print(
                f"[sweep] running sheet={sheet} mode={cfg['retrieval_mode']} alpha={cfg['hybrid_alpha']:.2f}"
            )
            scores = evaluate(
                args.queries,
                args.qrels,
                retrieve_shard=args.retrieve_shard,
                rerank_shard=args.rerank_shard,
                limit=int(args.limit),
                rerank_k=int(args.rerank_k),
                retrieval_mode=str(cfg["retrieval_mode"]),
                hybrid_alpha=float(cfg["hybrid_alpha"]),
                rerank_source=args.rerank_source,
                filter_missing_qrels=bool(args.filter_missing_qrels),
                per_query_csv=per_query_csv,
                per_query_topk_csv=None,
            )

            row = {
                "run": sheet,
                "retrieval_mode": cfg["retrieval_mode"],
                "hybrid_alpha": cfg["hybrid_alpha"],
            }
            for m in metric_headers:
                row[m] = float(scores.get(m, 0.0))
            run_outputs.append(row)
            per_query_by_run[sheet] = _read_csv_rows(per_query_csv)
    _build_master_matplotlib_plot(run_outputs, metric_headers, plot_png_path, style_path)

    wb = Workbook()
    master = wb.active
    master.title = "master"

    master_headers = ["run", "retrieval_mode", "hybrid_alpha", *metric_headers]
    _write_table(master, 1, master_headers, run_outputs)
    master.freeze_panes = "A2"
    _auto_fit_columns(master, max_col=len(master_headers), max_row=len(run_outputs) + 1)
    if plot_png_path is not None and plot_png_path.exists():
        img = XLImage(str(plot_png_path))
        master.add_image(img, "J2")
    else:
        master.cell(row=2, column=10, value="Plot generation failed: no image found.")

    for cfg in RUNS:
        sheet = cfg["sheet"]
        ws = wb.create_sheet(title=sheet[:31])
        summary_headers = ["metric", "value"]
        summary_rows = []
        run_row = next(r for r in run_outputs if r["run"] == sheet)
        for m in metric_headers:
            summary_rows.append({"metric": m, "value": run_row.get(m, 0.0)})

        row_cursor = _write_table(ws, 1, summary_headers, summary_rows)
        row_cursor += 2

        per_query_rows = per_query_by_run.get(sheet, [])
        if per_query_rows:
            pq_headers = list(per_query_rows[0].keys())
            row_cursor = _write_table(ws, row_cursor, pq_headers, per_query_rows)
            ws.freeze_panes = "A2"
            _auto_fit_columns(ws, max_col=len(pq_headers), max_row=row_cursor)
        else:
            ws.cell(row=row_cursor, column=1, value="No per-query rows generated.")
            _auto_fit_columns(ws, max_col=2, max_row=row_cursor)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.output)
    print(f"[done] wrote workbook: {args.output}")


if __name__ == "__main__":
    main()
