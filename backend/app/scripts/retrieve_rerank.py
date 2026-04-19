"""
Minimal retrieval + ColBERT-style reranking utilities.

- Retrieve Claim candidates from Weaviate using vector, BM25, or hybrid retrieval.
- Rerank with token vectors from LMDB or vectors fetched from Weaviate.
- Optionally evaluate simple IR metrics from query/qrels files.
- Provide a direct patent (doc_id) lookup helper.
"""
from __future__ import annotations

import argparse
import atexit
import csv
import json
import math
import os
import random
import re
import tempfile
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import numpy as np
import requests
from weaviate.collections.classes.filters import Filter

from backend.app.embed import embed_query_tokens_for_shard, model, normalize_text_for_embedding, tokenizer
from backend.app.services.download import rich_to_plain
from backend.app.store import (
    LMDB_PATH_128_F16,
    LMDB_PATH_128_F32,
    get_client,
    load_colbert_from_lmdb,
    resolve_lmdb_path,
)
from backend.app.vector_config import (
    WEAVIATE_NAMED_VECTOR,
    assert_128_variant,
)


def _env_bool(name: str, default: bool = False) -> bool:
    raw = os.environ.get(name)
    if raw is None:
        return default
    return str(raw).strip().lower() in {"1", "true", "yes", "on"}


_DEFAULT_HTTP_HOST = os.environ.get(
    "WEAVIATE_HTTP_HOST",
    os.environ.get("WEAVIATE_LOCAL_HOST", "127.0.0.1"),
).strip() or "127.0.0.1"
_DEFAULT_HTTP_PORT = int(
    os.environ.get(
        "WEAVIATE_HTTP_PORT",
        os.environ.get("WEAVIATE_LOCAL_PORT", "8081"),
    )
)
_DEFAULT_HTTP_SCHEME = "https" if _env_bool("WEAVIATE_HTTP_SECURE", False) else "http"
_DEFAULT_WEAVIATE_HTTP_BASE = f"{_DEFAULT_HTTP_SCHEME}://{_DEFAULT_HTTP_HOST}:{_DEFAULT_HTTP_PORT}"

WEAVIATE_GRAPHQL = os.environ.get("WEAVIATE_GRAPHQL", f"{_DEFAULT_WEAVIATE_HTTP_BASE}/v1/graphql")
WEAVIATE_OBJECTS = os.environ.get("WEAVIATE_OBJECTS", f"{_DEFAULT_WEAVIATE_HTTP_BASE}/v1/objects")
WEAVIATE_REQUEST_TIMEOUT_S = float(os.environ.get("WEAVIATE_REQUEST_TIMEOUT_S", "300"))
WEAVIATE_VECTOR_TIMEOUT_INITIAL_S = float(os.environ.get("WEAVIATE_VECTOR_TIMEOUT_INITIAL_S", "180"))
WEAVIATE_VECTOR_TIMEOUT_STEP_S = float(os.environ.get("WEAVIATE_VECTOR_TIMEOUT_STEP_S", "60"))
WEAVIATE_VECTOR_TIMEOUT_MAX_S = float(os.environ.get("WEAVIATE_VECTOR_TIMEOUT_MAX_S", "300"))


SHARD_TO_PATH = {
    "128_f32": LMDB_PATH_128_F32,
    "128_f16": LMDB_PATH_128_F16,
}
VALID_RETRIEVAL_MODES = ("vector", "bm25", "hybrid")
VALID_RERANK_SOURCES = ("lmdb", "weaviate")
WEAVIATE_VECTOR_FETCH_BATCH_SIZE = max(1, int(os.environ.get("WEAVIATE_VECTOR_FETCH_BATCH_SIZE", "128")))
WEAVIATE_VECTOR_FETCH_MODE = os.environ.get("WEAVIATE_VECTOR_FETCH_MODE", "auto").strip().lower() or "auto"

_WEAVIATE_VECTOR_CLIENT = None
_WEAVIATE_ENDPOINT_PRINTED = False


def _normalize_retrieval_mode(value: str | None) -> str:
    mode = str(value or "vector").strip().lower()
    if mode not in VALID_RETRIEVAL_MODES:
        raise ValueError(f"Unknown retrieval mode '{value}'. Choose from: {list(VALID_RETRIEVAL_MODES)}")
    return mode


def _safe_float(value: str | None, default: float) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def _safe_int(value: str | None, default: int) -> int:
    try:
        return int(value)
    except (TypeError, ValueError):
        return default


def _clamp_hybrid_alpha(value: float) -> float:
    return max(0.0, min(1.0, float(value)))


def _normalize_rerank_source(value: str | None) -> str:
    source = str(value or DEFAULT_RERANK_SOURCE).strip().lower()
    if source not in VALID_RERANK_SOURCES:
        raise ValueError(f"Unknown rerank source '{value}'. Choose from: {list(VALID_RERANK_SOURCES)}")
    return source


DEFAULT_RETRIEVAL_MODE = _normalize_retrieval_mode(
    os.environ.get("RETRIEVAL_MODE", os.environ.get("WEAVIATE_RETRIEVAL_MODE", "vector"))
)
DEFAULT_HYBRID_ALPHA = _clamp_hybrid_alpha(
    _safe_float(os.environ.get("HYBRID_ALPHA", os.environ.get("WEAVIATE_HYBRID_ALPHA", "0.5")), 0.5)
)
FORCE_CLIENT_HYBRID = os.environ.get("FORCE_CLIENT_HYBRID", "1").strip() not in {
    "0",
    "false",
    "False",
    "no",
    "NO",
}
DEFAULT_RERANK_SOURCE = os.environ.get("RERANK_SOURCE", "weaviate").strip().lower()
if DEFAULT_RERANK_SOURCE not in VALID_RERANK_SOURCES:
    raise ValueError(
        f"Invalid RERANK_SOURCE='{DEFAULT_RERANK_SOURCE}'. Allowed: {list(VALID_RERANK_SOURCES)}"
    )
DEBUG_COMPARE_SAMPLE_N = int(os.environ.get("DEBUG_VECTOR_COMPARE_N", "10"))
DEFAULT_PER_QUERY_TOPK = max(1, _safe_int(os.environ.get("PER_QUERY_TOPK"), 10))


@dataclass
class ClaimHit:
    uuid: str
    claim_id: str
    doc_id: str
    claim_type: str = ""
    text: str = ""
    distance: float | None = None
    score: float | None = None


def _is_vector_graphql_query(query: str) -> bool:
    query_lc = str(query or "").lower()
    return "nearvector" in query_lc or "hybrid:{" in query_lc


def _vector_timeout_schedule_s() -> list[float]:
    start_s = max(1.0, float(WEAVIATE_VECTOR_TIMEOUT_INITIAL_S))
    step_s = max(1.0, float(WEAVIATE_VECTOR_TIMEOUT_STEP_S))
    max_s = max(start_s, float(WEAVIATE_VECTOR_TIMEOUT_MAX_S))
    schedule: list[float] = []
    current = start_s
    while current <= max_s:
        schedule.append(current)
        current += step_s
    if not schedule:
        schedule.append(start_s)
    if schedule[-1] < max_s:
        schedule.append(max_s)
    return schedule


def _graphql_timeout_schedule_s(query: str) -> list[float]:
    if _is_vector_graphql_query(query):
        return _vector_timeout_schedule_s()
    return [max(1.0, float(WEAVIATE_REQUEST_TIMEOUT_S))]


def _is_retryable_graphql_payload_error(payload_errors: object) -> bool:
    text = str(payload_errors or "").lower()
    retry_markers = (
        "timeout",
        "timed out",
        "deadline",
        "temporarily unavailable",
        "connection reset",
        "connection refused",
    )
    return any(marker in text for marker in retry_markers)


def _post_graphql(query: str) -> dict:
    global _WEAVIATE_ENDPOINT_PRINTED
    if not _WEAVIATE_ENDPOINT_PRINTED:
        print(f"[weaviate] graphql={WEAVIATE_GRAPHQL}")
        _WEAVIATE_ENDPOINT_PRINTED = True
    timeout_schedule = _graphql_timeout_schedule_s(query)
    for attempt_idx, timeout_s in enumerate(timeout_schedule, start=1):
        is_last_attempt = attempt_idx >= len(timeout_schedule)
        try:
            resp = requests.post(WEAVIATE_GRAPHQL, json={"query": query}, timeout=timeout_s)
            if resp.status_code != 200:
                print("STATUS:", resp.status_code)
                print("RESPONSE:", resp.text)
                if resp.status_code == 404 and "/v1/graphql" in WEAVIATE_GRAPHQL:
                    print(
                        "[hint] 404 on /v1/graphql. If your Weaviate is on 8081, set "
                        "WEAVIATE_GRAPHQL=http://localhost:8081/v1/graphql and "
                        "WEAVIATE_OBJECTS=http://localhost:8081/v1/objects "
                        "(or unset both and set WEAVIATE_HTTP_PORT=8081)."
                    )
                if resp.status_code in {429, 500, 502, 503, 504} and not is_last_attempt:
                    print(
                        "[warn] GraphQL HTTP %s on attempt %s/%s; retrying with timeout=%.0fs"
                        % (resp.status_code, attempt_idx, len(timeout_schedule), timeout_schedule[attempt_idx])
                    )
                    continue
            resp.raise_for_status()
            payload = resp.json()
            if payload.get("errors"):
                if not is_last_attempt and _is_retryable_graphql_payload_error(payload["errors"]):
                    print(
                        "[warn] GraphQL payload error on attempt %s/%s; retrying with timeout=%.0fs"
                        % (attempt_idx, len(timeout_schedule), timeout_schedule[attempt_idx])
                    )
                    continue
                raise RuntimeError(payload["errors"])
            return payload["data"]
        except requests.exceptions.Timeout as exc:
            if not is_last_attempt:
                print(
                    "[warn] GraphQL timeout on attempt %s/%s at %.0fs; retrying with %.0fs"
                    % (attempt_idx, len(timeout_schedule), timeout_s, timeout_schedule[attempt_idx])
                )
                continue
            raise TimeoutError(
                "Weaviate GraphQL request timed out after %.0fs (attempt %s/%s)."
                % (timeout_s, attempt_idx, len(timeout_schedule))
            ) from exc

    raise RuntimeError("Weaviate GraphQL request failed without returning data.")


def _embed_query_dense(query: str) -> np.ndarray:
    # Mean-pooled dense vector for Weaviate nearVector retrieval.
    import torch

    query_text = normalize_text_for_embedding(query)
    tokens = tokenizer(
        query_text,
        return_tensors="pt",
        truncation=True,
        padding="max_length",
        max_length=256,
    )
    tokens = {k: v.to(model.device) for k, v in tokens.items()}
    model.eval()
    with torch.no_grad():
        outputs = model(**tokens)
    token_embeddings = outputs.last_hidden_state.squeeze(0)
    attn_mask = tokens["attention_mask"].squeeze(0).bool()
    masked = token_embeddings[attn_mask]
    if masked.numel() == 0:
        masked = token_embeddings[:1]
    dense = masked.mean(dim=0, keepdim=True).detach().cpu().numpy().astype("float32")
    return dense[0]


def _embed_query_tokens(query: str, shard: str) -> np.ndarray:
    """
    Token-level query embeddings matching the LMDB shard dimensionality.
    """
    shard = assert_128_variant(shard, context="retrieve_rerank._embed_query_tokens")
    variants = embed_query_tokens_for_shard(query, shard=shard, max_length=256)
    arr = np.asarray(variants)
    if arr.ndim == 1:
        arr = arr.reshape(1, -1)
    if arr.ndim != 2 or arr.shape[1] != 128:
        raise ValueError(f"Query token vectors must be [T,128], got {arr.shape}.")
    return arr


def _build_claim_retrieval_args(
    query: str,
    limit: int,
    *,
    shard: str,
    retrieval_mode: str,
    hybrid_alpha: float,
) -> str:
    mode = _normalize_retrieval_mode(retrieval_mode)
    if mode == "bm25":
        return f"bm25: {{query: {json.dumps(query)}}}, limit: {int(limit)}"
    if mode == "hybrid":
        alpha = _clamp_hybrid_alpha(hybrid_alpha)
        # Avoid Weaviate edge-case instability on hybrid alpha=0 by using native BM25.
        if alpha <= 0.0:
            return f"bm25: {{query: {json.dumps(query)}}}, limit: {int(limit)}"
        mv_vecs = _embed_query_tokens(query, shard).tolist()
        if alpha >= 1.0:
            return (
                f'nearVector: {{vector: {mv_vecs}, targetVectors: ["{WEAVIATE_NAMED_VECTOR}"]}}, '
                f"limit: {int(limit)}"
            )
        return (
            "hybrid: {"
            f"query: {json.dumps(query)}, "
            f"alpha: {alpha:.6f}, "
            f"vector: {mv_vecs}, "
            f'targetVectors: ["{WEAVIATE_NAMED_VECTOR}"]'
            f"}}, limit: {int(limit)}"
        )
    mv_vecs = _embed_query_tokens(query, shard).tolist()
    return (
        f'nearVector: {{vector: {mv_vecs}, targetVectors: ["{WEAVIATE_NAMED_VECTOR}"]}}, '
        f"limit: {int(limit)}"
    )


def _query_claim_rows(retrieval_args: str) -> list[dict]:
    gql = (
        "{ Get { Claim("
        f"{retrieval_args}"
        ") { claim_id doc_id claim_type text _additional { id distance } } } }"
    )
    data = _post_graphql(gql)
    return data.get("Get", {}).get("Claim", []) or []


def _coerce_token_matrix(value: object, *, identity: str) -> np.ndarray:
    arr = np.asarray(value, dtype=np.float32)
    if arr.ndim == 1:
        arr = arr.reshape(1, -1)
    if arr.ndim != 2 or arr.shape[0] <= 0:
        raise ValueError(f"Weaviate vector for {identity} is empty or malformed: shape={arr.shape}")
    if arr.shape[1] != 128:
        raise ValueError(f"Weaviate vector for {identity} must be 128-d, got shape={arr.shape}.")
    return arr


def _get_weaviate_vector_client():
    global _WEAVIATE_VECTOR_CLIENT
    if _WEAVIATE_VECTOR_CLIENT is None:
        _WEAVIATE_VECTOR_CLIENT = get_client()

        def _close_client() -> None:
            global _WEAVIATE_VECTOR_CLIENT
            client = _WEAVIATE_VECTOR_CLIENT
            _WEAVIATE_VECTOR_CLIENT = None
            if client is not None:
                try:
                    client.close()
                except Exception:
                    pass

        atexit.register(_close_client)
    return _WEAVIATE_VECTOR_CLIENT


def _chunked(values: list[str], size: int) -> Iterable[list[str]]:
    for start in range(0, len(values), max(1, int(size))):
        yield values[start:start + max(1, int(size))]


def _fetch_colbert_vectors_from_weaviate_batched(object_ids: list[str]) -> dict[str, np.ndarray]:
    ids = [str(oid).strip() for oid in object_ids if str(oid).strip()]
    if not ids:
        return {}
    client = _get_weaviate_vector_client()
    collection = client.collections.get("Claim")
    out: dict[str, np.ndarray] = {}
    for chunk in _chunked(ids, WEAVIATE_VECTOR_FETCH_BATCH_SIZE):
        response = collection.query.fetch_objects(
            limit=len(chunk),
            filters=Filter.by_id().contains_any(chunk),
            include_vector=[WEAVIATE_NAMED_VECTOR],
            return_properties=False,
        )
        for obj in response.objects:
            oid = str(getattr(obj, "uuid", "") or "").strip()
            vectors = getattr(obj, "vector", None)
            if not oid:
                continue
            if not isinstance(vectors, dict):
                raise ValueError(
                    f"Weaviate object {oid} missing named vectors map; expected dict with key '{WEAVIATE_NAMED_VECTOR}'."
                )
            if WEAVIATE_NAMED_VECTOR not in vectors:
                raise ValueError(
                    f"Weaviate object {oid} missing named vector '{WEAVIATE_NAMED_VECTOR}'. "
                    f"Available keys={sorted(vectors.keys())}"
                )
            out[oid] = _coerce_token_matrix(vectors[WEAVIATE_NAMED_VECTOR], identity=oid)

        missing = [oid for oid in chunk if oid not in out]
        if missing:
            raise RuntimeError(
                f"Batched Weaviate vector fetch returned {len(chunk) - len(missing)}/{len(chunk)} objects. "
                f"Missing ids: {missing[:5]}{'...' if len(missing) > 5 else ''}"
            )
    return out


def _fetch_colbert_vectors_from_weaviate_http(object_ids: list[str]) -> dict[str, np.ndarray]:
    """
    Fetch token-level ColBERT vectors from Weaviate object REST API.
    Endpoint used: GET /v1/objects/Claim/{uuid}?include=vector
    """
    ids = [str(oid).strip() for oid in object_ids if str(oid).strip()]
    if not ids:
        return {}
    sess = requests.Session()
    out: dict[str, np.ndarray] = {}
    for oid in ids:
        url = f"{WEAVIATE_OBJECTS}/Claim/{oid}"
        resp = sess.get(url, params={"include": "vector"}, timeout=WEAVIATE_REQUEST_TIMEOUT_S)
        if resp.status_code != 200:
            raise RuntimeError(f"Weaviate vector fetch failed for {oid}: {resp.status_code} {resp.text[:300]}")
        payload = resp.json()
        vectors = payload.get("vectors")
        if not isinstance(vectors, dict):
            raise ValueError(
                f"Weaviate object {oid} missing named vectors map; expected dict with key '{WEAVIATE_NAMED_VECTOR}'."
            )
        if WEAVIATE_NAMED_VECTOR not in vectors:
            raise ValueError(
                f"Weaviate object {oid} missing named vector '{WEAVIATE_NAMED_VECTOR}'. "
                f"Available keys={sorted(vectors.keys())}"
            )
        out[oid] = _coerce_token_matrix(vectors[WEAVIATE_NAMED_VECTOR], identity=oid)

    missing = [oid for oid in ids if oid not in out]
    if missing:
        raise RuntimeError(f"Failed to fetch vectors for object ids: {missing}")
    return out


def fetch_colbert_vectors_from_weaviate(object_ids: list[str]) -> dict[str, np.ndarray]:
    ids = [str(oid).strip() for oid in object_ids if str(oid).strip()]
    if not ids:
        return {}
    mode = WEAVIATE_VECTOR_FETCH_MODE
    if mode not in {"auto", "batch", "http"}:
        raise ValueError(
            f"Invalid WEAVIATE_VECTOR_FETCH_MODE='{mode}'. Allowed=['auto', 'batch', 'http']"
        )
    if mode in {"auto", "batch"}:
        try:
            return _fetch_colbert_vectors_from_weaviate_batched(ids)
        except Exception as exc:
            if mode == "batch":
                raise
            print(f"[warn] batched Weaviate vector fetch failed; falling back to per-object HTTP fetch: {exc}")
    return _fetch_colbert_vectors_from_weaviate_http(ids)


def _rows_to_hits(items: list[dict]) -> list[ClaimHit]:
    hits: list[ClaimHit] = []
    for it in items:
        addl = it.get("_additional") or {}
        hits.append(
            ClaimHit(
                uuid=addl.get("id", ""),
                distance=addl.get("distance"),
                claim_id=it.get("claim_id", ""),
                doc_id=it.get("doc_id", ""),
                claim_type=str(it.get("claim_type") or ""),
                text=rich_to_plain(str(it.get("text") or "")),
            )
        )
    return hits


def _claim_key(hit: ClaimHit) -> str:
    return str(hit.claim_id or hit.uuid or hit.doc_id or "")


_CLAIM_BASE_RE = re.compile(r"^(?P<prefix>.+-CLM-)(?P<num>\d+)(?P<rest>.*)$", re.IGNORECASE)


def _normalize_claim_id(value: str) -> str:
    return str(value or "").strip()


def _claim_base_id(claim_id: str) -> str:
    """
    Canonicalize potentially chunked claim IDs to a base claim id.
    Examples:
    - DOC-CLM-1 -> DOC-CLM-1
    - DOC-CLM-1-CHUNK-2 -> DOC-CLM-1
    - DOC-CLM-1_part2 -> DOC-CLM-1
    """
    cid = _normalize_claim_id(claim_id)
    if not cid:
        return ""
    match = _CLAIM_BASE_RE.match(cid)
    if not match:
        return cid
    prefix = match.group("prefix")
    num = match.group("num")
    rest = match.group("rest") or ""
    # Keep canonicalization conservative: only collapse when suffix does not continue digits.
    if rest and rest[:1].isdigit():
        return cid
    return f"{prefix}{num}"


def _dedupe_preserve_order(values: list[str]) -> list[str]:
    out: list[str] = []
    seen: set[str] = set()
    for v in values:
        if not v or v in seen:
            continue
        seen.add(v)
        out.append(v)
    return out


def _is_relevant_claim_id(claim_id: str, relevant_base_ids: set[str]) -> bool:
    base = _claim_base_id(claim_id)
    return bool(base) and base in relevant_base_ids


def _fuse_hybrid_hits(
    vector_hits: list[ClaimHit],
    bm25_hits: list[ClaimHit],
    alpha: float,
    limit: int,
    *,
    rrf_k: int = 60,
) -> list[ClaimHit]:
    alpha = _clamp_hybrid_alpha(alpha)
    weights = {
        "vector": alpha,
        "bm25": 1.0 - alpha,
    }
    score_by_key: dict[str, float] = {}
    hit_by_key: dict[str, ClaimHit] = {}

    for source, hits in (("vector", vector_hits), ("bm25", bm25_hits)):
        w = weights[source]
        if w <= 0.0:
            continue
        for rank, hit in enumerate(hits, start=1):
            key = _claim_key(hit)
            if not key:
                continue
            if key not in hit_by_key:
                hit_by_key[key] = hit
            else:
                if hit_by_key[key].distance is None and hit.distance is not None:
                    hit_by_key[key].distance = hit.distance
                if not hit_by_key[key].doc_id and hit.doc_id:
                    hit_by_key[key].doc_id = hit.doc_id
                if not hit_by_key[key].claim_type and hit.claim_type:
                    hit_by_key[key].claim_type = hit.claim_type
                if not hit_by_key[key].text and hit.text:
                    hit_by_key[key].text = hit.text
            score_by_key[key] = score_by_key.get(key, 0.0) + (w / float(rrf_k + rank))

    ranked_keys = sorted(score_by_key, key=lambda k: score_by_key[k], reverse=True)[: int(limit)]
    return [hit_by_key[k] for k in ranked_keys]


def _retrieve_hybrid_client_fusion(
    query: str,
    limit: int,
    shard: str,
    alpha: float,
) -> list[ClaimHit]:
    bm25_args = f"bm25: {{query: {json.dumps(query)}}}, limit: {int(limit)}"
    vec_args = (
        f'nearVector: {{vector: {_embed_query_tokens(query, shard).tolist()}, '
        f'targetVectors: ["{WEAVIATE_NAMED_VECTOR}"]}}, limit: {int(limit)}'
    )
    vector_items: list[dict] = []
    bm25_items: list[dict] = []
    try:
        vector_items = _query_claim_rows(vec_args)
    except Exception as ve:
        print(f"[warn] hybrid nearVector leg failed: {ve}")
    try:
        bm25_items = _query_claim_rows(bm25_args)
    except Exception as be:
        print(f"[warn] hybrid bm25 leg failed: {be}")
    if not vector_items and not bm25_items:
        raise RuntimeError("both hybrid legs failed (nearVector and bm25)")
    return _fuse_hybrid_hits(
        _rows_to_hits(vector_items),
        _rows_to_hits(bm25_items),
        alpha,
        int(limit),
    )


def retrieve_claims(
    query: str,
    limit: int = 200,
    shard: str = "128_f16",
    *,
    retrieval_mode: str = DEFAULT_RETRIEVAL_MODE,
    hybrid_alpha: float = DEFAULT_HYBRID_ALPHA,
) -> list[ClaimHit]:
    shard = assert_128_variant(shard, context="retrieve_claims")
    mode = _normalize_retrieval_mode(retrieval_mode)
    if mode in {"vector", "hybrid"} and shard not in SHARD_TO_PATH:
        raise ValueError(f"Unknown shard '{shard}'. Choose from: {sorted(SHARD_TO_PATH)}")
    alpha = _clamp_hybrid_alpha(hybrid_alpha)
    retrieval_args = _build_claim_retrieval_args(
        query,
        limit,
        shard=shard,
        retrieval_mode=mode,
        hybrid_alpha=alpha,
    )

    if mode == "hybrid" and 0.0 < alpha < 1.0 and FORCE_CLIENT_HYBRID:
        hits = _retrieve_hybrid_client_fusion(query, int(limit), shard, alpha)
        return hits

    try:
        hits = _rows_to_hits(_query_claim_rows(retrieval_args))
    except Exception as e:
        if mode != "hybrid" or alpha <= 0.0 or alpha >= 1.0:
            raise
        print(f"[warn] server-side hybrid failed; using client-side fusion fallback: {e}")
        hits = _retrieve_hybrid_client_fusion(query, int(limit), shard, alpha)

    return hits


def _maxsim_score(query_tokens: np.ndarray, doc_tokens: np.ndarray) -> float:
    if doc_tokens.size == 0 or query_tokens.size == 0:
        return float("-inf")
    qt = query_tokens.astype(np.float32, copy=False)
    dt = doc_tokens.astype(np.float32, copy=False)
    sims = qt @ dt.T
    return float(np.max(sims, axis=1).sum())


def _fmt_score(value: float | None) -> str:
    return "" if value is None else f"{value:.6f}"


def rerank_with_lmdb(
    hits: list[ClaimHit],
    query: str,
    shard: str,
    rerank_k: int = 100,
    query_tokens: np.ndarray | None = None,
) -> list[ClaimHit]:
    shard = assert_128_variant(shard, context="rerank_with_lmdb")
    q_tokens = query_tokens if query_tokens is not None else _embed_query_tokens(query, shard)
    top = hits[: int(rerank_k)]
    rescored: list[ClaimHit] = []
    for h in top:
        lmdb_path = resolve_lmdb_path(shard, doc_id=h.doc_id or None)
        doc_tokens = load_colbert_from_lmdb(lmdb_path, h.claim_id)
        if doc_tokens is None:
            raise RuntimeError(f"Missing LMDB vectors for claim_id={h.claim_id} at {lmdb_path}")
        if np.asarray(doc_tokens).ndim != 2 or np.asarray(doc_tokens).shape[1] != 128:
            raise ValueError(
                f"LMDB vectors for claim_id={h.claim_id} must be [T,128], got {np.asarray(doc_tokens).shape}"
            )
        h.score = _maxsim_score(q_tokens, doc_tokens)
        rescored.append(h)
    rescored.sort(key=lambda x: (x.score if x.score is not None else float("-inf")), reverse=True)
    return rescored + hits[int(rerank_k) :]


def rerank_with_weaviate(
    hits: list[ClaimHit],
    query: str,
    shard: str,
    rerank_k: int = 100,
    query_tokens: np.ndarray | None = None,
) -> list[ClaimHit]:
    shard = assert_128_variant(shard, context="rerank_with_weaviate")
    q_tokens = query_tokens if query_tokens is not None else _embed_query_tokens(query, shard)
    top = hits[: int(rerank_k)]
    object_ids = [h.uuid for h in top]
    if any(not oid for oid in object_ids):
        raise ValueError("All rerank candidates must include Weaviate object UUID for rerank_source=weaviate.")
    vectors_by_object_id = fetch_colbert_vectors_from_weaviate(object_ids)
    rescored: list[ClaimHit] = []
    for h in top:
        doc_tokens = vectors_by_object_id.get(h.uuid)
        if doc_tokens is None:
            raise RuntimeError(f"Missing Weaviate vectors for object id={h.uuid}")
        if doc_tokens.ndim != 2 or doc_tokens.shape[1] != 128:
            raise ValueError(f"Weaviate vectors for object id={h.uuid} must be [T,128], got {doc_tokens.shape}")
        h.score = _maxsim_score(q_tokens, doc_tokens)
        rescored.append(h)
    rescored.sort(key=lambda x: (x.score if x.score is not None else float("-inf")), reverse=True)
    return rescored + hits[int(rerank_k) :]


def rerank_hits(
    hits: list[ClaimHit],
    query: str,
    shard: str,
    rerank_k: int = 100,
    *,
    rerank_source: str = DEFAULT_RERANK_SOURCE,
    query_tokens: np.ndarray | None = None,
) -> list[ClaimHit]:
    source = _normalize_rerank_source(rerank_source)
    if source == "lmdb":
        return rerank_with_lmdb(hits, query, shard=shard, rerank_k=rerank_k, query_tokens=query_tokens)
    return rerank_with_weaviate(hits, query, shard=shard, rerank_k=rerank_k, query_tokens=query_tokens)


def debug_compare_lmdb_vs_weaviate_vectors(
    hits: list[ClaimHit],
    shard: str,
    sample_n: int,
    *,
    max_abs_tol: float = 0.05,
) -> dict[str, float]:
    shard = assert_128_variant(shard, context="debug_compare_lmdb_vs_weaviate_vectors")
    sample_pool = [h for h in hits if h.uuid and h.claim_id]
    if not sample_pool:
        raise ValueError("No candidates with both uuid and claim_id available for vector comparison.")
    n = min(int(sample_n), len(sample_pool))
    sampled = random.sample(sample_pool, n)
    obj_ids = [h.uuid for h in sampled]
    weaviate_vecs = fetch_colbert_vectors_from_weaviate(obj_ids)
    diffs: list[float] = []
    for h in sampled:
        lmdb_path = resolve_lmdb_path(shard, doc_id=h.doc_id or None)
        lmdb_vec = load_colbert_from_lmdb(lmdb_path, h.claim_id)
        if lmdb_vec is None:
            raise RuntimeError(f"LMDB vector missing for claim_id={h.claim_id}")
        wv_vec = weaviate_vecs.get(h.uuid)
        if wv_vec is None:
            raise RuntimeError(f"Weaviate vector missing for object_id={h.uuid}")
        lmdb_arr = np.asarray(lmdb_vec, dtype=np.float32)
        if lmdb_arr.ndim == 1:
            lmdb_arr = lmdb_arr.reshape(1, -1)
        if lmdb_arr.shape != wv_vec.shape:
            raise ValueError(
                f"Shape mismatch claim_id={h.claim_id}: lmdb={lmdb_arr.shape} weaviate={wv_vec.shape}"
            )
        max_abs = float(np.max(np.abs(lmdb_arr - wv_vec)))
        diffs.append(max_abs)
    stats = {
        "n": float(n),
        "max_abs_min": float(np.min(diffs)),
        "max_abs_mean": float(np.mean(diffs)),
        "max_abs_p95": float(np.percentile(diffs, 95)),
        "max_abs_max": float(np.max(diffs)),
    }
    print(
        "[debug-compare] n={n} max_abs(min/mean/p95/max)="
        "{minv:.6f}/{meanv:.6f}/{p95:.6f}/{maxv:.6f}".format(
            n=int(stats["n"]),
            minv=stats["max_abs_min"],
            meanv=stats["max_abs_mean"],
            p95=stats["max_abs_p95"],
            maxv=stats["max_abs_max"],
        )
    )
    if stats["max_abs_max"] > float(max_abs_tol):
        raise ValueError(
            f"LMDB vs Weaviate vector mismatch too large: max_abs={stats['max_abs_max']:.6f} > {max_abs_tol}"
        )
    return stats


def lookup_patent(doc_id: str, limit: int = 200) -> list[ClaimHit]:
    where = (
        "{path:[\"doc_id\"],operator:Equal,valueText:"
        f"\"{doc_id}\"}}"
    )
    gql = (
        "{ Get { Claim("
        f"where:{where}, limit:{int(limit)}"
        ") { claim_id doc_id claim_type text _additional { id } } } }"
    )
    data = _post_graphql(gql)
    return _rows_to_hits(data["Get"]["Claim"])


def _lookup_claim_object_ids_by_claim_id(
    claim_id: str,
    *,
    include_chunk_variants: bool = True,
    like_limit: int = 256,
) -> list[str]:
    """
    Resolve Weaviate object UUIDs for a claim_id.
    If include_chunk_variants=True, expands by claim_id prefix and filters rows
    back to the same canonical base claim id to support chunked IDs.
    """
    raw_claim_id = _normalize_claim_id(claim_id)
    if not raw_claim_id:
        return []
    base_claim_id = _claim_base_id(raw_claim_id)
    if not base_claim_id:
        return []

    if include_chunk_variants:
        where = (
            "{path:[\"claim_id\"],operator:Like,valueText:"
            f"{json.dumps(base_claim_id + '*')}}}"
        )
        limit_clause = f", limit:{int(max(1, like_limit))}"
    else:
        where = (
            "{path:[\"claim_id\"],operator:Equal,valueText:"
            f"{json.dumps(base_claim_id)}}}"
        )
        limit_clause = ", limit:1"

    gql = (
        "{ Get { Claim("
        f"where:{where}{limit_clause}"
        ") { claim_id _additional { id } } } }"
    )
    data = _post_graphql(gql)
    rows = data.get("Get", {}).get("Claim", []) or []

    out: list[str] = []
    seen: set[str] = set()
    for row in rows:
        row_claim_id = _normalize_claim_id(row.get("claim_id"))
        if _claim_base_id(row_claim_id) != base_claim_id:
            continue
        addl = row.get("_additional") or {}
        object_id = _normalize_claim_id(addl.get("id"))
        if not object_id or object_id in seen:
            continue
        seen.add(object_id)
        out.append(object_id)
    return out


def _lookup_claim_uuid_by_claim_id(claim_id: str) -> str | None:
    object_ids = _lookup_claim_object_ids_by_claim_id(claim_id, include_chunk_variants=True)
    return object_ids[0] if object_ids else None


def _claim_id_exists(claim_id: str) -> bool:
    return _lookup_claim_uuid_by_claim_id(claim_id) is not None


def precision_at_k(ranked_ids: list[str], relevant: set[str], k: int) -> float:
    if k <= 0:
        return 0.0
    top = ranked_ids[:k]
    if not top:
        return 0.0
    hits = sum(1 for rid in top if rid in relevant)
    return hits / float(len(top))


def recall_at_k(ranked_ids: list[str], relevant: set[str], k: int) -> float:
    if not relevant:
        return 0.0
    top = ranked_ids[:k]
    hits = sum(1 for rid in top if rid in relevant)
    return hits / float(len(relevant))


def ndcg_at_k(ranked_ids: list[str], relevant: set[str], k: int) -> float:
    def dcg(ids: Iterable[str]) -> float:
        total = 0.0
        for i, rid in enumerate(ids, start=1):
            rel = 1.0 if rid in relevant else 0.0
            if rel > 0:
                total += rel / math.log2(i + 1.0)
        return total

    top = ranked_ids[:k]
    if not top:
        return 0.0
    ideal = [1.0] * min(len(relevant), k)
    idcg = sum(rel / math.log2(i + 1.0) for i, rel in enumerate(ideal, start=1))
    if idcg == 0:
        return 0.0
    return dcg(top) / idcg


def mrr_at_k(ranked_ids: list[str], relevant: set[str], k: int) -> float:
    for i, rid in enumerate(ranked_ids[:k], start=1):
        if rid in relevant:
            return 1.0 / float(i)
    return 0.0


def _first_relevant_rank(ranked_ids: list[str], relevant: set[str]) -> int | None:
    for i, rid in enumerate(ranked_ids, start=1):
        if rid in relevant:
            return i
    return None


def _write_per_query_csv(rows: list[dict], path: Path, *, topk: int = 10) -> None:
    topk = max(1, int(topk))
    path.parent.mkdir(parents=True, exist_ok=True)
    relevant_claim_cols = [f"relevant_claim_{i}" for i in range(1, topk + 1)]
    relevant_claim_maxsim_cols = [f"relevant_claim_{i}_maxsim" for i in range(1, topk + 1)]
    relevant_claim_candidate_rank_cols = [f"relevant_claim_{i}_candidate_rank" for i in range(1, topk + 1)]
    relevant_claim_reranked_rank_cols = [f"relevant_claim_{i}_reranked_rank" for i in range(1, topk + 1)]
    retrieved_claim_id_cols = [f"retrieved_top{i}_claim_id" for i in range(1, topk + 1)]
    retrieved_claim_maxsim_cols = [f"retrieved_top{i}_maxsim" for i in range(1, topk + 1)]
    retrieved_claim_rel_cols = [f"retrieved_top{i}_is_relevant" for i in range(1, topk + 1)]
    retrieved_claim_text_cols = [f"retrieved_top{i}_claim_text" for i in range(1, topk + 1)]
    reranked_claim_id_cols = [f"reranked_top{i}_claim_id" for i in range(1, topk + 1)]
    reranked_claim_maxsim_cols = [f"reranked_top{i}_maxsim" for i in range(1, topk + 1)]
    reranked_claim_text_cols = [f"reranked_top{i}_claim_text" for i in range(1, topk + 1)]
    reranked_claim_rel_cols = [f"reranked_top{i}_is_relevant" for i in range(1, topk + 1)]
    fieldnames = [
        "status",
        "query",
        "num_relevant",
        "relevant_claim_extra_count",
        *relevant_claim_cols,
        *relevant_claim_maxsim_cols,
        *relevant_claim_candidate_rank_cols,
        *relevant_claim_reranked_rank_cols,
        "candidate_k",
        "rerank_k",
        "candidate_hit",
        "candidate_hits_count",
        "candidate_first_relevant_rank",
        "candidate_relevant_ranks",
        "reranked_hit_at_10",
        "reranked_hits_at_10",
        "reranked_first_relevant_rank",
        "reranked_relevant_ranks_at_10",
        "precision@10",
        "recall@10",
        "ndcg@10",
        "mrr@10",
        *retrieved_claim_id_cols,
        *retrieved_claim_maxsim_cols,
        *retrieved_claim_rel_cols,
        *retrieved_claim_text_cols,
        *reranked_claim_id_cols,
        *reranked_claim_maxsim_cols,
        *reranked_claim_rel_cols,
        *reranked_claim_text_cols,
    ]
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def _write_per_query_topk_csv(rows: list[dict], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = [
        "status",
        "query",
        "num_relevant",
        "relevant_claim_ids",
        "candidate_k",
        "rerank_k",
        "rank",
        "candidate_rank",
        "reranked_rank",
        "claim_id",
        "doc_id",
        "claim_type",
        "candidate_maxsim",
        "reranked_maxsim",
        "score",
        "candidate_distance",
        "distance",
        "is_relevant",
        "text",
    ]
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def _write_per_query_xlsx(
    per_query_rows: list[dict],
    per_query_topk_rows: list[dict],
    path: Path,
    *,
    topk: int = 10,
) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment
    except ImportError as exc:
        raise RuntimeError(
            "openpyxl is required for --per-query-xlsx. Install with: "
            "python -m pip install openpyxl"
        ) from exc

    topk = max(1, int(topk))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    metrics_ws = wb.active
    metrics_ws.title = "metrics"

    text_cols = {f"reranked_top{i}_claim_text" for i in range(1, topk + 1)}
    metric_columns: list[str] = []
    for row in per_query_rows:
        for key in row.keys():
            if key in text_cols:
                continue
            if key not in metric_columns:
                metric_columns.append(key)
    if metric_columns:
        metrics_ws.append(metric_columns)
        for row in per_query_rows:
            metrics_ws.append([row.get(c, "") for c in metric_columns])
        metrics_ws.freeze_panes = "A2"

    normalized_claim_rows = per_query_topk_rows
    if not normalized_claim_rows and per_query_rows:
        normalized_claim_rows = []
        for row in per_query_rows:
            query = str(row.get("query") or "")
            status = str(row.get("status") or "")
            num_relevant = row.get("num_relevant", "")
            candidate_k = row.get("candidate_k", "")
            rerank_k = row.get("rerank_k", "")
            relevant_ids = []
            for i in range(1, topk + 1):
                rel_id = str(row.get(f"relevant_claim_{i}", "") or "").strip()
                if rel_id:
                    relevant_ids.append(rel_id)
            rel_joined = "|".join(relevant_ids)
            for rank in range(1, topk + 1):
                normalized_claim_rows.append(
                    {
                        "status": status,
                        "query": query,
                        "num_relevant": num_relevant,
                        "relevant_claim_ids": rel_joined,
                        "candidate_k": candidate_k,
                        "rerank_k": rerank_k,
                        "rank": rank,
                        "candidate_rank": "",
                        "reranked_rank": rank,
                        "claim_id": row.get(f"reranked_top{rank}_claim_id", ""),
                        "doc_id": "",
                        "claim_type": "",
                        "candidate_maxsim": "",
                        "reranked_maxsim": row.get(f"reranked_top{rank}_maxsim", ""),
                        "score": row.get(f"reranked_top{rank}_maxsim", ""),
                        "candidate_distance": "",
                        "distance": "",
                        "is_relevant": row.get(f"reranked_top{rank}_is_relevant", False),
                        "text": row.get(f"reranked_top{rank}_claim_text", ""),
                    }
                )

    query_order: list[str] = []
    seen_queries: set[str] = set()
    for row in per_query_rows:
        q = str(row.get("query") or "").strip()
        if q and q not in seen_queries:
            seen_queries.add(q)
            query_order.append(q)
    for row in normalized_claim_rows:
        q = str(row.get("query") or "").strip()
        if q and q not in seen_queries:
            seen_queries.add(q)
            query_order.append(q)

    per_query_status: dict[str, str] = {}
    for row in per_query_rows:
        q = str(row.get("query") or "").strip()
        if not q:
            continue
        per_query_status[q] = str(row.get("status") or "")

    grouped_claim_rows: dict[str, list[dict]] = {q: [] for q in query_order}
    for row in normalized_claim_rows:
        q = str(row.get("query") or "").strip()
        if not q:
            continue
        grouped_claim_rows.setdefault(q, []).append(row)

    def _rank_key(value: object) -> int:
        try:
            return int(value)
        except Exception:
            return 10**9

    def _safe_sheet_name(query: str, idx: int, used: set[str]) -> str:
        invalid = set(r'[]:*?/\\')
        cleaned = "".join(" " if ch in invalid else ch for ch in query).replace("'", "").strip()
        cleaned = " ".join(cleaned.split())
        base = cleaned or f"query_{idx:02d}"
        prefix = f"Q{idx:02d}_"
        max_base_len = max(1, 31 - len(prefix))
        base = base[:max_base_len]
        candidate = f"{prefix}{base}"
        if candidate not in used:
            used.add(candidate)
            return candidate
        n = 2
        while True:
            suffix = f"_{n}"
            max_len = max(1, 31 - len(prefix) - len(suffix))
            candidate = f"{prefix}{base[:max_len]}{suffix}"
            if candidate not in used:
                used.add(candidate)
                return candidate
            n += 1

    used_sheet_names: set[str] = {"metrics"}
    query_columns = [
        "rank",
        "candidate_rank",
        "reranked_rank",
        "claim_id",
        "doc_id",
        "claim_type",
        "candidate_maxsim",
        "reranked_maxsim",
        "score",
        "candidate_distance",
        "distance",
        "is_relevant",
        "claim_text",
    ]
    for idx, query in enumerate(query_order, start=1):
        ws = wb.create_sheet(title=_safe_sheet_name(query, idx, used_sheet_names))
        ws["A1"] = "query"
        ws["B1"] = query
        ws["A2"] = "status"
        ws["B2"] = per_query_status.get(query, "")
        ws.append([])
        ws.append(query_columns)

        rows_for_query = sorted(grouped_claim_rows.get(query, []), key=lambda r: _rank_key(r.get("rank")))
        rows_for_query = [
            r
            for r in rows_for_query
            if str(r.get("claim_id") or "").strip() or str(r.get("text") or "").strip()
        ]
        if not rows_for_query:
            ws.append(["", "", "", "", "", "", "", "", "", "", "", "", "No reranked claims found."])
        else:
            for row in rows_for_query:
                ws.append(
                    [
                        row.get("rank", ""),
                        row.get("candidate_rank", ""),
                        row.get("reranked_rank", ""),
                        row.get("claim_id", ""),
                        row.get("doc_id", ""),
                        row.get("claim_type", ""),
                        row.get("candidate_maxsim", ""),
                        row.get("reranked_maxsim", ""),
                        row.get("score", ""),
                        row.get("candidate_distance", ""),
                        row.get("distance", ""),
                        row.get("is_relevant", False),
                        row.get("text", ""),
                    ]
                )

    def _autosize(ws, *, max_width: int = 80) -> None:
        for col in ws.columns:
            letter = col[0].column_letter
            width = 10
            for cell in col:
                value = "" if cell.value is None else str(cell.value)
                width = max(width, min(max_width, len(value) + 2))
            ws.column_dimensions[letter].width = width

    _autosize(metrics_ws, max_width=64)
    for ws in wb.worksheets:
        if ws.title == "metrics":
            continue
        _autosize(ws, max_width=60)
        ws.column_dimensions["B"].width = 110 if ws["A1"].value == "query" else ws.column_dimensions["B"].width
        text_col_idx = None
        for i, cell in enumerate(ws[4], start=1):
            if str(cell.value or "") == "claim_text":
                text_col_idx = i
                break
        if text_col_idx is not None:
            text_col_letter = ws.cell(row=4, column=text_col_idx).column_letter
            ws.column_dimensions[text_col_letter].width = 110
            for row in ws.iter_rows(min_row=5, min_col=text_col_idx, max_col=text_col_idx):
                row[0].alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A5"

    wb.save(path)


def _write_queries_from_qrels(qrels_path: Path, queries_path: Path) -> int:
    seen: set[str] = set()
    queries_path.parent.mkdir(parents=True, exist_ok=True)
    with qrels_path.open("r", encoding="utf-8") as src, queries_path.open("w", encoding="utf-8") as dst:
        for line in src:
            line = line.strip()
            if not line:
                continue
            row = json.loads(line)
            q = str(row.get("query") or "").strip()
            if not q or q in seen:
                continue
            seen.add(q)
            dst.write(json.dumps({"query": q}, ensure_ascii=False))
            dst.write("\n")
    return len(seen)


def evaluate(
    queries_path: Path,
    qrels_path: Path,
    retrieve_shard: str,
    rerank_shard: str,
    limit: int,
    rerank_k: int,
    *,
    retrieval_mode: str = DEFAULT_RETRIEVAL_MODE,
    hybrid_alpha: float = DEFAULT_HYBRID_ALPHA,
    rerank_source: str = DEFAULT_RERANK_SOURCE,
    debug_compare_vectors: bool = False,
    debug_compare_n: int = DEBUG_COMPARE_SAMPLE_N,
    debug_compare_max_abs_tol: float = 0.05,
    filter_missing_qrels: bool = False,
    per_query_csv: Path | None = None,
    per_query_topk_csv: Path | None = None,
    per_query_xlsx: Path | None = None,
    per_query_topk: int = DEFAULT_PER_QUERY_TOPK,
) -> dict:
    print(f"[eval] filter_missing_qrels={bool(filter_missing_qrels)}")
    retrieval_mode = _normalize_retrieval_mode(retrieval_mode)
    hybrid_alpha = _clamp_hybrid_alpha(hybrid_alpha)
    rerank_source = _normalize_rerank_source(rerank_source)
    rerank_shard = assert_128_variant(rerank_shard, context="evaluate.rerank_shard")
    retrieve_shard = assert_128_variant(retrieve_shard, context="evaluate.retrieve_shard")
    topk = max(1, int(per_query_topk))
    with queries_path.open("r", encoding="utf-8") as f:
        queries = [json.loads(line) for line in f if line.strip()]
    # Normalize query text early so query->qrels matching is not affected by whitespace.
    for row in queries:
        row["query"] = str(row.get("query") or "").strip()
    with qrels_path.open("r", encoding="utf-8") as f:
        qrels: dict[str, set[str]] = {}
        for row in (json.loads(line) for line in f if line.strip()):
            q = str(row.get("query") or "").strip()
            if not q:
                continue
            rel_ids = {
                str(claim_id).strip()
                for claim_id in (row.get("relevant_claim_ids", []) or [])
                if str(claim_id).strip()
            }
            if q in qrels:
                qrels[q].update(rel_ids)
            else:
                qrels[q] = rel_ids

    candidate_hit_metric = f"candidate_hit@{int(limit)}"
    metrics = {
        "precision@10": [],
        "recall@10": [],
        "ndcg@10": [],
        "mrr@10": [],
        candidate_hit_metric: [],
    }
    per_query_rows: list[dict] = []
    per_query_topk_rows: list[dict] = []
    need_per_query_rows = per_query_csv is not None or per_query_xlsx is not None
    need_per_query_topk_rows = per_query_topk_csv is not None or per_query_xlsx is not None
    total_q = len(queries)
    if filter_missing_qrels:
        cache: dict[str, bool] = {}
        filtered = []
        dropped = 0
        for row in queries:
            q = row.get("query", "")
            rel = qrels.get(q, set())
            keep = False
            for claim_id in rel:
                if claim_id not in cache:
                    cache[claim_id] = _claim_id_exists(claim_id)
                if cache[claim_id]:
                    keep = True
                    break
            if keep:
                filtered.append(row)
            else:
                dropped += 1
                if need_per_query_rows:
                    sorted_rel = sorted(rel)
                    row_out = {
                        "status": "filtered_missing_qrels",
                        "query": q,
                        "num_relevant": len(rel),
                        "relevant_claim_extra_count": max(0, len(sorted_rel) - topk),
                        "candidate_k": int(limit),
                        "rerank_k": int(rerank_k),
                        "candidate_hit": False,
                        "candidate_hits_count": 0,
                        "candidate_first_relevant_rank": "",
                        "candidate_relevant_ranks": "",
                        "reranked_hit_at_10": False,
                        "reranked_hits_at_10": 0,
                        "reranked_first_relevant_rank": "",
                        "reranked_relevant_ranks_at_10": "",
                        "precision@10": 0.0,
                        "recall@10": 0.0,
                        "ndcg@10": 0.0,
                        "mrr@10": 0.0,
                    }
                    for idx in range(1, topk + 1):
                        row_out[f"relevant_claim_{idx}"] = sorted_rel[idx - 1] if idx <= len(sorted_rel) else ""
                        row_out[f"relevant_claim_{idx}_maxsim"] = ""
                        row_out[f"relevant_claim_{idx}_candidate_rank"] = ""
                        row_out[f"relevant_claim_{idx}_reranked_rank"] = ""
                        row_out[f"retrieved_top{idx}_claim_id"] = ""
                        row_out[f"retrieved_top{idx}_maxsim"] = ""
                        row_out[f"retrieved_top{idx}_is_relevant"] = False
                        row_out[f"retrieved_top{idx}_claim_text"] = ""
                        row_out[f"reranked_top{idx}_claim_id"] = ""
                        row_out[f"reranked_top{idx}_maxsim"] = ""
                        row_out[f"reranked_top{idx}_is_relevant"] = False
                        row_out[f"reranked_top{idx}_claim_text"] = ""
                    per_query_rows.append(row_out)
                if need_per_query_topk_rows:
                    for rank in range(1, topk + 1):
                        sorted_rel = sorted(rel)
                        per_query_topk_rows.append(
                            {
                                "status": "filtered_missing_qrels",
                                "query": q,
                                "num_relevant": len(rel),
                                "relevant_claim_ids": "|".join(sorted(rel)),
                                "candidate_k": int(limit),
                                "rerank_k": int(rerank_k),
                                "rank": rank,
                                "candidate_rank": "",
                                "reranked_rank": rank,
                                "claim_id": "",
                                "doc_id": "",
                                "claim_type": "",
                                "candidate_maxsim": "",
                                "reranked_maxsim": "",
                                "score": "",
                                "candidate_distance": "",
                                "distance": "",
                                "is_relevant": False,
                                "text": "",
                            }
                        )
        queries = filtered
        total_q = len(queries)
        print(f"[eval] filtered {dropped} query(ies) with no relevant claims in index")
    for i, row in enumerate(queries, start=1):
        q = row.get("query", "")
        if not q:
            continue
        print(f"[eval] {i}/{total_q} retrieving + reranking")
        query_start_s = time.perf_counter()
        retrieval_start_s = time.perf_counter()
        rel = qrels.get(q, set())
        hits = retrieve_claims(
            q,
            limit=limit,
            shard=retrieve_shard,
            retrieval_mode=retrieval_mode,
            hybrid_alpha=hybrid_alpha,
        )
        retrieval_ms = (time.perf_counter() - retrieval_start_s) * 1000.0
        print(
            f"[timing] q={i}/{total_q} retrieval_ms={retrieval_ms:.1f} candidates={len(hits)}"
        )
        candidate_ids = [h.claim_id for h in hits]
        candidate_base_ids = [_claim_base_id(cid) for cid in candidate_ids]
        rel_base_ids = {
            _claim_base_id(claim_id)
            for claim_id in rel
            if _claim_base_id(claim_id)
        }
        candidate_rank_by_base: dict[str, int] = {}
        for idx, base_id in enumerate(candidate_base_ids, start=1):
            if not base_id or base_id in candidate_rank_by_base:
                continue
            candidate_rank_by_base[base_id] = idx
        candidate_hit_by_base: dict[str, ClaimHit] = {}
        for hit in hits:
            base_id = _claim_base_id(hit.claim_id)
            if base_id and base_id not in candidate_hit_by_base:
                candidate_hit_by_base[base_id] = hit
        candidate_relevant_rank_by_base = {
            base_id: rank
            for base_id, rank in candidate_rank_by_base.items()
            if base_id in rel_base_ids
        }
        candidate_ranks = sorted(candidate_relevant_rank_by_base.values())
        candidate_hit = bool(candidate_ranks)
        q_tokens = _embed_query_tokens(q, rerank_shard)
        candidate_object_ids_by_base: dict[str, list[str]] = {}
        for hit in hits:
            base_id = _claim_base_id(hit.claim_id)
            object_id = _normalize_claim_id(hit.uuid)
            if not base_id or not object_id:
                continue
            bucket = candidate_object_ids_by_base.setdefault(base_id, [])
            if object_id not in bucket:
                bucket.append(object_id)
        claim_object_ids_lookup_cache: dict[str, list[str]] = {}
        weaviate_vec_cache: dict[str, np.ndarray] = {}
        object_maxsim_cache: dict[str, float | None] = {}
        maxsim_cache: dict[str, float | None] = {}

        def _resolve_claim_object_ids(claim_id: str) -> list[str]:
            base_id = _claim_base_id(claim_id)
            if not base_id:
                return []
            if base_id in claim_object_ids_lookup_cache:
                return claim_object_ids_lookup_cache[base_id]
            candidate_ids_for_base = candidate_object_ids_by_base.get(base_id) or []
            if candidate_ids_for_base:
                claim_object_ids_lookup_cache[base_id] = list(candidate_ids_for_base)
                return claim_object_ids_lookup_cache[base_id]
            resolved_ids = _lookup_claim_object_ids_by_claim_id(base_id, include_chunk_variants=True)
            claim_object_ids_lookup_cache[base_id] = resolved_ids
            return resolved_ids

        def _object_maxsim(object_id: str, *, claim_id: str = "") -> float | None:
            oid = _normalize_claim_id(object_id)
            if not oid:
                return None
            cached = object_maxsim_cache.get(oid)
            if cached is not None or oid in object_maxsim_cache:
                return cached
            if rerank_source != "weaviate":
                object_maxsim_cache[oid] = None
                return None
            if oid not in weaviate_vec_cache:
                fetched = fetch_colbert_vectors_from_weaviate([oid])
                weaviate_vec_cache.update(fetched)
            doc_tokens = weaviate_vec_cache.get(oid)
            if doc_tokens is None:
                raise RuntimeError(
                    f"Missing Weaviate vectors for object_id={oid} claim_id={claim_id or 'unknown'}"
                )
            arr = np.asarray(doc_tokens)
            if arr.ndim == 1:
                arr = arr.reshape(1, -1)
            if arr.ndim != 2 or arr.shape[1] != 128:
                raise ValueError(
                    f"MaxSim vectors must be [T,128], got shape={arr.shape} for object_id={oid} "
                    f"claim_id={claim_id or 'unknown'}"
                )
            score = _maxsim_score(q_tokens, doc_tokens)
            object_maxsim_cache[oid] = score
            return score

        def _claim_maxsim(claim_id: str) -> float | None:
            base_id = _claim_base_id(claim_id)
            if not base_id:
                return None
            cached = maxsim_cache.get(base_id)
            if cached is not None or base_id in maxsim_cache:
                return cached
            if rerank_source == "weaviate":
                object_ids = _resolve_claim_object_ids(base_id)
                if not object_ids:
                    maxsim_cache[base_id] = None
                    return None
                scores = [
                    _object_maxsim(object_id, claim_id=base_id)
                    for object_id in object_ids
                ]
                scores = [score for score in scores if score is not None]
                if not scores:
                    maxsim_cache[base_id] = None
                    return None
                score = max(scores)
                maxsim_cache[base_id] = score
                return score
            else:
                lmdb_path = resolve_lmdb_path(rerank_shard)
                doc_tokens = load_colbert_from_lmdb(lmdb_path, base_id)
                if doc_tokens is None:
                    maxsim_cache[base_id] = None
                    return None
            arr = np.asarray(doc_tokens)
            if arr.ndim == 1:
                arr = arr.reshape(1, -1)
            if arr.ndim != 2 or arr.shape[1] != 128:
                raise ValueError(
                    f"MaxSim vectors must be [T,128], got shape={arr.shape} for claim_id={base_id}"
                )
            score = _maxsim_score(q_tokens, doc_tokens)
            maxsim_cache[base_id] = score
            return score

        rerank_start_s = time.perf_counter()
        ranked = rerank_hits(
            hits,
            q,
            shard=rerank_shard,
            rerank_k=rerank_k,
            rerank_source=rerank_source,
            query_tokens=q_tokens,
        )
        rerank_ms = (time.perf_counter() - rerank_start_s) * 1000.0
        total_ms = (time.perf_counter() - query_start_s) * 1000.0
        print(
            f"[timing] q={i}/{total_q} rerank_ms={rerank_ms:.1f} total_ms={total_ms:.1f}"
        )
        if debug_compare_vectors and rerank_source == "weaviate":
            debug_compare_lmdb_vs_weaviate_vectors(
                ranked[: max(1, rerank_k)],
                rerank_shard,
                sample_n=debug_compare_n,
                max_abs_tol=debug_compare_max_abs_tol,
            )
        ranked_ids = [h.claim_id for h in ranked]
        ranked_base_ids = [_claim_base_id(cid) for cid in ranked_ids]
        reranked_rank_by_base: dict[str, int] = {}
        for idx, base_id in enumerate(ranked_base_ids, start=1):
            if not base_id or base_id in reranked_rank_by_base:
                continue
            reranked_rank_by_base[base_id] = idx
        ranked_metric_ids = _dedupe_preserve_order(ranked_base_ids)
        p10 = precision_at_k(ranked_metric_ids, rel_base_ids, 10)
        r10 = recall_at_k(ranked_metric_ids, rel_base_ids, 10)
        n10 = ndcg_at_k(ranked_metric_ids, rel_base_ids, 10)
        m10 = mrr_at_k(ranked_metric_ids, rel_base_ids, 10)
        metrics["precision@10"].append(p10)
        metrics["recall@10"].append(r10)
        metrics["ndcg@10"].append(n10)
        metrics["mrr@10"].append(m10)
        metrics[candidate_hit_metric].append(1.0 if candidate_hit else 0.0)

        if need_per_query_rows:
            sorted_rel = sorted(rel)
            retrieved_topk = hits[:topk]
            reranked_topk = ranked[:topk]
            reranked_top10 = ranked[:10]
            reranked_top10_rank_by_base: dict[str, int] = {}
            for idx, hit in enumerate(reranked_top10, start=1):
                base_id = _claim_base_id(hit.claim_id)
                if not base_id or base_id not in rel_base_ids or base_id in reranked_top10_rank_by_base:
                    continue
                reranked_top10_rank_by_base[base_id] = idx
            reranked_top10_ranks = sorted(reranked_top10_rank_by_base.values())
            row_out = {
                "status": "evaluated",
                "query": q,
                "num_relevant": len(rel),
                "relevant_claim_extra_count": max(0, len(sorted_rel) - topk),
                "candidate_k": int(limit),
                "rerank_k": int(rerank_k),
                "candidate_hit": candidate_hit,
                "candidate_hits_count": len(candidate_ranks),
                "candidate_first_relevant_rank": candidate_ranks[0] if candidate_ranks else "",
                "candidate_relevant_ranks": "|".join(str(r) for r in candidate_ranks),
                "reranked_hit_at_10": bool(reranked_top10_rank_by_base),
                "reranked_hits_at_10": len(reranked_top10_rank_by_base),
                "reranked_first_relevant_rank": _first_relevant_rank(ranked_base_ids, rel_base_ids) or "",
                "reranked_relevant_ranks_at_10": "|".join(str(r) for r in reranked_top10_ranks),
                "precision@10": p10,
                "recall@10": r10,
                "ndcg@10": n10,
                "mrr@10": m10,
            }
            for idx in range(1, topk + 1):
                row_out[f"relevant_claim_{idx}"] = sorted_rel[idx - 1] if idx <= len(sorted_rel) else ""
                rel_claim = sorted_rel[idx - 1] if idx <= len(sorted_rel) else ""
                rel_claim_base = _claim_base_id(rel_claim)
                row_out[f"relevant_claim_{idx}_maxsim"] = _fmt_score(_claim_maxsim(rel_claim))
                row_out[f"relevant_claim_{idx}_candidate_rank"] = candidate_rank_by_base.get(rel_claim_base, "")
                row_out[f"relevant_claim_{idx}_reranked_rank"] = reranked_rank_by_base.get(rel_claim_base, "")
                if idx <= len(retrieved_topk):
                    cand = retrieved_topk[idx - 1]
                    row_out[f"retrieved_top{idx}_claim_id"] = cand.claim_id
                    candidate_score = _object_maxsim(cand.uuid, claim_id=cand.claim_id) if cand.uuid else None
                    if candidate_score is None:
                        candidate_score = _claim_maxsim(cand.claim_id)
                    row_out[f"retrieved_top{idx}_maxsim"] = _fmt_score(candidate_score)
                    row_out[f"retrieved_top{idx}_is_relevant"] = _is_relevant_claim_id(cand.claim_id, rel_base_ids)
                    row_out[f"retrieved_top{idx}_claim_text"] = cand.text
                else:
                    row_out[f"retrieved_top{idx}_claim_id"] = ""
                    row_out[f"retrieved_top{idx}_maxsim"] = ""
                    row_out[f"retrieved_top{idx}_is_relevant"] = False
                    row_out[f"retrieved_top{idx}_claim_text"] = ""
                if idx <= len(reranked_topk):
                    hit = reranked_topk[idx - 1]
                    row_out[f"reranked_top{idx}_claim_id"] = hit.claim_id
                    hit_score = hit.score if hit.score is not None else _claim_maxsim(hit.claim_id)
                    row_out[f"reranked_top{idx}_maxsim"] = _fmt_score(hit_score)
                    row_out[f"reranked_top{idx}_is_relevant"] = _is_relevant_claim_id(hit.claim_id, rel_base_ids)
                    row_out[f"reranked_top{idx}_claim_text"] = hit.text
                else:
                    row_out[f"reranked_top{idx}_claim_id"] = ""
                    row_out[f"reranked_top{idx}_maxsim"] = ""
                    row_out[f"reranked_top{idx}_is_relevant"] = False
                    row_out[f"reranked_top{idx}_claim_text"] = ""
            per_query_rows.append(row_out)
        if need_per_query_topk_rows:
            sorted_rel = "|".join(sorted(rel))
            topk_hits = ranked[:topk]
            for rank, h in enumerate(topk_hits, start=1):
                candidate_base_id = _claim_base_id(h.claim_id)
                candidate_hit_row = candidate_hit_by_base.get(candidate_base_id)
                candidate_rank = candidate_rank_by_base.get(candidate_base_id, "")
                candidate_score_value = (
                    _object_maxsim(candidate_hit_row.uuid, claim_id=h.claim_id)
                    if candidate_hit_row is not None and candidate_hit_row.uuid
                    else None
                )
                if candidate_score_value is None:
                    candidate_score_value = _claim_maxsim(h.claim_id)
                candidate_score = _fmt_score(candidate_score_value)
                reranked_score = _fmt_score(h.score if h.score is not None else _claim_maxsim(h.claim_id))
                per_query_topk_rows.append(
                    {
                        "status": "evaluated",
                        "query": q,
                        "num_relevant": len(rel),
                        "relevant_claim_ids": sorted_rel,
                        "candidate_k": int(limit),
                        "rerank_k": int(rerank_k),
                        "rank": rank,
                        "candidate_rank": candidate_rank,
                        "reranked_rank": rank,
                        "claim_id": h.claim_id,
                        "doc_id": h.doc_id,
                        "claim_type": h.claim_type,
                        "candidate_maxsim": candidate_score,
                        "reranked_maxsim": reranked_score,
                        "score": reranked_score,
                        "candidate_distance": (
                            "" if candidate_hit_row is None or candidate_hit_row.distance is None
                            else f"{candidate_hit_row.distance:.6f}"
                        ),
                        "distance": "" if h.distance is None else f"{h.distance:.6f}",
                        "is_relevant": _is_relevant_claim_id(h.claim_id, rel_base_ids),
                        "text": h.text,
                    }
                )
            if len(topk_hits) < topk:
                for rank in range(len(topk_hits) + 1, topk + 1):
                    per_query_topk_rows.append(
                        {
                            "status": "evaluated",
                            "query": q,
                            "num_relevant": len(rel),
                            "relevant_claim_ids": sorted_rel,
                            "candidate_k": int(limit),
                            "rerank_k": int(rerank_k),
                            "rank": rank,
                            "candidate_rank": "",
                            "reranked_rank": rank,
                            "claim_id": "",
                            "doc_id": "",
                            "claim_type": "",
                            "candidate_maxsim": "",
                            "reranked_maxsim": "",
                            "score": "",
                            "candidate_distance": "",
                            "distance": "",
                            "is_relevant": False,
                            "text": "",
                        }
                    )

    if per_query_csv is not None:
        _write_per_query_csv(per_query_rows, per_query_csv, topk=topk)
        print(f"Wrote per-query diagnostics to {per_query_csv}")
    if per_query_topk_csv is not None:
        _write_per_query_topk_csv(per_query_topk_rows, per_query_topk_csv)
        print(f"Wrote per-query top-{topk} reranked claims to {per_query_topk_csv}")
    if per_query_xlsx is not None:
        _write_per_query_xlsx(per_query_rows, per_query_topk_rows, per_query_xlsx, topk=topk)
        print(f"Wrote per-query workbook to {per_query_xlsx}")

    return {k: (sum(v) / len(v) if v else 0.0) for k, v in metrics.items()}


def evaluate_sweep(
    queries_path: Path,
    qrels_path: Path,
    retrieve_shard: str,
    rerank_shards: list[str],
    limit: int,
    rerank_k: int,
    *,
    retrieval_mode: str = DEFAULT_RETRIEVAL_MODE,
    hybrid_alpha: float = DEFAULT_HYBRID_ALPHA,
    rerank_source: str = DEFAULT_RERANK_SOURCE,
    filter_missing_qrels: bool = False,
) -> list[dict]:
    results: list[dict] = []
    for shard in rerank_shards:
        scores = evaluate(
            queries_path,
            qrels_path,
            retrieve_shard=retrieve_shard,
            rerank_shard=shard,
            limit=limit,
            rerank_k=rerank_k,
            retrieval_mode=retrieval_mode,
            hybrid_alpha=hybrid_alpha,
            rerank_source=rerank_source,
            filter_missing_qrels=filter_missing_qrels,
        )
        row = {"rerank_shard": shard, **scores}
        results.append(row)
    return results


def _build_alpha_sweep(start: float, stop: float, step: float) -> list[float]:
    if step <= 0:
        raise ValueError("--alpha-step must be > 0.")
    if start < 0.0 or start > 1.0 or stop < 0.0 or stop > 1.0:
        raise ValueError("--alpha-start and --alpha-stop must be in [0, 1].")
    if stop < start:
        raise ValueError("--alpha-stop must be >= --alpha-start.")

    values: list[float] = []
    cursor = float(start)
    eps = max(1e-9, abs(step) * 1e-9)
    while cursor <= float(stop) + eps:
        values.append(round(cursor, 6))
        cursor += float(step)

    if not values:
        values = [round(float(start), 6)]

    if abs(values[-1] - float(stop)) > eps:
        values.append(round(float(stop), 6))

    deduped: list[float] = []
    seen: set[float] = set()
    for a in values:
        if a not in seen:
            seen.add(a)
            deduped.append(a)
    return deduped


def evaluate_hybrid_alpha_sweep(
    queries_path: Path,
    qrels_path: Path,
    retrieve_shard: str,
    rerank_shard: str,
    limit: int,
    rerank_k: int,
    hybrid_alphas: list[float],
    *,
    rerank_source: str = DEFAULT_RERANK_SOURCE,
    filter_missing_qrels: bool = False,
) -> list[dict]:
    results: list[dict] = []
    for alpha in hybrid_alphas:
        scores = evaluate(
            queries_path,
            qrels_path,
            retrieve_shard=retrieve_shard,
            rerank_shard=rerank_shard,
            limit=limit,
            rerank_k=rerank_k,
            retrieval_mode="hybrid",
            hybrid_alpha=float(alpha),
            rerank_source=rerank_source,
            filter_missing_qrels=filter_missing_qrels,
        )
        row = {"hybrid_alpha": float(alpha), **scores}
        results.append(row)
    return results


def _print_hits(hits: list[ClaimHit], k: int = 10):
    for i, h in enumerate(hits[:k], start=1):
        score = f"{h.score:.3f}" if h.score is not None else "-"
        dist = f"{h.distance:.4f}" if h.distance is not None else "-"
        print(f"{i:02d} score={score} dist={dist} claim_id={h.claim_id} doc_id={h.doc_id} type={h.claim_type}")
        print(f"    {h.text[:200]}")


def _write_hits_csv(hits: list[ClaimHit], path: Path, k: int = 10):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["rank", "score", "distance", "claim_id", "doc_id", "claim_type", "text"])
        for i, h in enumerate(hits[:k], start=1):
            writer.writerow(
                [
                    i,
                    "" if h.score is None else f"{h.score:.6f}",
                    "" if h.distance is None else f"{h.distance:.6f}",
                    h.claim_id,
                    h.doc_id,
                    h.claim_type,
                    h.text,
                ]
            )


def main():
    parser = argparse.ArgumentParser(description="Retrieve + rerank claims with shard selection.")
    parser.add_argument("--query", type=str, default="", help="Query text for retrieval.")
    parser.add_argument("--limit", type=int, default=200, help="Initial retrieval limit.")
    parser.add_argument("--rerank-k", type=int, default=100, help="How many initial hits to rerank.")
    parser.add_argument(
        "--retrieval-mode",
        type=str,
        default=DEFAULT_RETRIEVAL_MODE,
        choices=list(VALID_RETRIEVAL_MODES),
        help="First-stage retrieval mode in Weaviate.",
    )
    parser.add_argument(
        "--hybrid-alpha",
        type=float,
        default=DEFAULT_HYBRID_ALPHA,
        help="Hybrid interpolation weight (0=bm25 only, 1=vector only).",
    )
    parser.add_argument("--save-before-csv", type=Path, default=None, help="Save pre-rerank top-k to CSV.")
    parser.add_argument("--save-after-csv", type=Path, default=None, help="Save reranked top-k to CSV.")
    parser.add_argument("--save-k", type=int, default=10, help="How many rows to write to CSV outputs.")
    parser.add_argument(
        "--retrieve-shard",
        type=str,
        default=os.environ.get("WEAVIATE_SHARD", "128_f16"),
        choices=sorted(SHARD_TO_PATH.keys()),
        help="Which shard to use for Weaviate retrieval (must match index dimension).",
    )
    parser.add_argument(
        "--rerank-shard",
        type=str,
        default=os.environ.get("COLBERT_SHARD", "128_f16"),
        choices=sorted(SHARD_TO_PATH.keys()),
        help="Which ColBERT shard to use for reranking.",
    )
    parser.add_argument(
        "--rerank-source",
        type=str,
        default=DEFAULT_RERANK_SOURCE,
        choices=list(VALID_RERANK_SOURCES),
        help="Source for token vectors during reranking.",
    )
    parser.add_argument(
        "--debug-compare-vectors",
        action=argparse.BooleanOptionalAction,
        default=False,
        help="When rerank-source=weaviate, compare fetched vectors against LMDB vectors.",
    )
    parser.add_argument(
        "--debug-compare-n",
        type=int,
        default=DEBUG_COMPARE_SAMPLE_N,
        help="Sample size for LMDB vs Weaviate vector comparison.",
    )
    parser.add_argument(
        "--debug-compare-max-abs-tol",
        type=float,
        default=0.05,
        help="Max allowed absolute difference for LMDB vs Weaviate debug comparison.",
    )
    parser.add_argument("--doc-id", type=str, default="", help="Lookup claims by patent doc_id.")
    parser.add_argument(
        "--queries",
        type=Path,
        default=None,
        help="Optional path to JSONL queries for evaluation. If omitted with --qrels, queries are derived from qrels.",
    )
    parser.add_argument(
        "--qrels",
        type=Path,
        default=None,
        help="Path to JSONL qrels for evaluation (rows may include query text).",
    )
    parser.add_argument("--metrics-csv", type=Path, default=None, help="Write evaluation metrics to CSV.")
    parser.add_argument(
        "--per-query-csv",
        type=Path,
        default=None,
        help="Write per-query diagnostics CSV (coverage, relevant ranks, and top-k IDs).",
    )
    parser.add_argument(
        "--per-query-topk-csv",
        "--per-query-top10-csv",
        dest="per_query_topk_csv",
        type=Path,
        default=None,
        help="Write one row per query per reranked rank (1-N) with claim metadata.",
    )
    parser.add_argument(
        "--per-query-topk",
        type=int,
        default=DEFAULT_PER_QUERY_TOPK,
        help="How many reranked rows to include per query when writing per-query CSV outputs.",
    )
    parser.add_argument(
        "--per-query-xlsx",
        type=Path,
        default=None,
        help="Write readable XLSX with one metrics sheet plus one reranked-claims sheet per query.",
    )
    parser.add_argument(
        "--sweep-rerank-shards",
        action="store_true",
        help="Run evaluation once per rerank shard and write a CSV.",
    )
    parser.add_argument(
        "--sweep-hybrid-alphas",
        action="store_true",
        help="Run evaluation across a hybrid-alpha range and write a CSV.",
    )
    parser.add_argument(
        "--alpha-start",
        type=float,
        default=0.0,
        help="Hybrid-alpha sweep start value (inclusive).",
    )
    parser.add_argument(
        "--alpha-stop",
        type=float,
        default=1.0,
        help="Hybrid-alpha sweep stop value (inclusive).",
    )
    parser.add_argument(
        "--alpha-step",
        type=float,
        default=0.2,
        help="Hybrid-alpha sweep step size.",
    )
    parser.add_argument(
        "--sweep-csv",
        type=Path,
        default=None,
        help="Output CSV path for sweep results (required with --sweep-rerank-shards or --sweep-hybrid-alphas).",
    )
    parser.add_argument(
        "--filter-missing-qrels",
        action=argparse.BooleanOptionalAction,
        default=_env_bool("FILTER_MISSING_QRELS", False),
        help="Drop queries whose relevant claim_ids are not present in Weaviate (default: off).",
    )
    args = parser.parse_args()
    args.hybrid_alpha = _clamp_hybrid_alpha(args.hybrid_alpha)
    args.rerank_source = _normalize_rerank_source(args.rerank_source)
    args.retrieve_shard = assert_128_variant(args.retrieve_shard, context="main.retrieve_shard")
    args.rerank_shard = assert_128_variant(args.rerank_shard, context="main.rerank_shard")
    args.per_query_topk = max(1, int(args.per_query_topk))
    print(f"[config] filter_missing_qrels={args.filter_missing_qrels}")
    temp_eval_dir = None
    eval_queries_path = args.queries
    if args.qrels and eval_queries_path is None:
        temp_eval_dir = tempfile.TemporaryDirectory(prefix="qrels_eval_")
        atexit.register(temp_eval_dir.cleanup)
        eval_queries_path = Path(temp_eval_dir.name) / "queries.jsonl"
        derived_count = _write_queries_from_qrels(args.qrels, eval_queries_path)
        print(f"[eval] derived {derived_count} quer{'y' if derived_count == 1 else 'ies'} from {args.qrels}")

    if args.doc_id:
        hits = lookup_patent(args.doc_id)
        print(f"Found {len(hits)} claims for doc_id={args.doc_id}")
        _print_hits(hits, k=min(10, len(hits)))
        if temp_eval_dir is not None:
            temp_eval_dir.cleanup()
        return

    if eval_queries_path and args.qrels:
        candidate_hit_metric = f"candidate_hit@{int(args.limit)}"
        if args.sweep_rerank_shards and args.sweep_hybrid_alphas:
            raise SystemExit("Choose one sweep mode: --sweep-rerank-shards or --sweep-hybrid-alphas.")
        if args.sweep_rerank_shards:
            if not args.sweep_csv:
                raise SystemExit("Provide --sweep-csv when using --sweep-rerank-shards.")
            if args.per_query_csv:
                raise SystemExit("--per-query-csv is not supported with --sweep-rerank-shards.")
            if args.per_query_topk_csv:
                raise SystemExit("--per-query-topk-csv is not supported with --sweep-rerank-shards.")
            if args.per_query_xlsx:
                raise SystemExit("--per-query-xlsx is not supported with --sweep-rerank-shards.")
            rerank_shards = sorted(SHARD_TO_PATH.keys())
            results = evaluate_sweep(
                eval_queries_path,
                args.qrels,
                retrieve_shard=args.retrieve_shard,
                rerank_shards=rerank_shards,
                limit=args.limit,
                rerank_k=args.rerank_k,
                retrieval_mode=args.retrieval_mode,
                hybrid_alpha=args.hybrid_alpha,
                rerank_source=args.rerank_source,
                filter_missing_qrels=args.filter_missing_qrels,
            )
            args.sweep_csv.parent.mkdir(parents=True, exist_ok=True)
            with args.sweep_csv.open("w", encoding="utf-8", newline="") as f:
                writer = csv.writer(f)
                writer.writerow(
                    [
                        "rerank_shard",
                        "precision@10",
                        "recall@10",
                        "ndcg@10",
                        "mrr@10",
                        candidate_hit_metric,
                    ]
                )
                for row in results:
                    writer.writerow(
                        [
                            row.get("rerank_shard"),
                            row.get("precision@10", 0.0),
                            row.get("recall@10", 0.0),
                            row.get("ndcg@10", 0.0),
                            row.get("mrr@10", 0.0),
                            row.get(candidate_hit_metric, 0.0),
                        ]
                    )
            print(f"Wrote sweep results to {args.sweep_csv}")
        elif args.sweep_hybrid_alphas:
            if not args.sweep_csv:
                raise SystemExit("Provide --sweep-csv when using --sweep-hybrid-alphas.")
            if args.per_query_csv:
                raise SystemExit("--per-query-csv is not supported with --sweep-hybrid-alphas.")
            if args.per_query_topk_csv:
                raise SystemExit("--per-query-topk-csv is not supported with --sweep-hybrid-alphas.")
            if args.per_query_xlsx:
                raise SystemExit("--per-query-xlsx is not supported with --sweep-hybrid-alphas.")
            try:
                alphas = _build_alpha_sweep(args.alpha_start, args.alpha_stop, args.alpha_step)
            except ValueError as exc:
                raise SystemExit(str(exc)) from exc
            results = evaluate_hybrid_alpha_sweep(
                eval_queries_path,
                args.qrels,
                retrieve_shard=args.retrieve_shard,
                rerank_shard=args.rerank_shard,
                limit=args.limit,
                rerank_k=args.rerank_k,
                hybrid_alphas=alphas,
                rerank_source=args.rerank_source,
                filter_missing_qrels=args.filter_missing_qrels,
            )
            args.sweep_csv.parent.mkdir(parents=True, exist_ok=True)
            with args.sweep_csv.open("w", encoding="utf-8", newline="") as f:
                writer = csv.writer(f)
                writer.writerow(
                    [
                        "hybrid_alpha",
                        "precision@10",
                        "recall@10",
                        "ndcg@10",
                        "mrr@10",
                        candidate_hit_metric,
                    ]
                )
                for row in results:
                    writer.writerow(
                        [
                            row.get("hybrid_alpha"),
                            row.get("precision@10", 0.0),
                            row.get("recall@10", 0.0),
                            row.get("ndcg@10", 0.0),
                            row.get("mrr@10", 0.0),
                            row.get(candidate_hit_metric, 0.0),
                        ]
                    )
            print(f"Wrote hybrid-alpha sweep results to {args.sweep_csv}")
        else:
            scores = evaluate(
                eval_queries_path,
                args.qrels,
                retrieve_shard=args.retrieve_shard,
                rerank_shard=args.rerank_shard,
                limit=args.limit,
                rerank_k=args.rerank_k,
                retrieval_mode=args.retrieval_mode,
                hybrid_alpha=args.hybrid_alpha,
                rerank_source=args.rerank_source,
                debug_compare_vectors=args.debug_compare_vectors,
                debug_compare_n=args.debug_compare_n,
                debug_compare_max_abs_tol=args.debug_compare_max_abs_tol,
                filter_missing_qrels=args.filter_missing_qrels,
                per_query_csv=args.per_query_csv,
                per_query_topk_csv=args.per_query_topk_csv,
                per_query_xlsx=args.per_query_xlsx,
                per_query_topk=args.per_query_topk,
            )
            print(json.dumps(scores, indent=2))
            if args.metrics_csv:
                args.metrics_csv.parent.mkdir(parents=True, exist_ok=True)
                with args.metrics_csv.open("w", encoding="utf-8", newline="") as f:
                    writer = csv.writer(f)
                    writer.writerow(["precision@10", "recall@10", "ndcg@10", "mrr@10", candidate_hit_metric])
                    writer.writerow(
                        [
                            scores.get("precision@10", 0.0),
                            scores.get("recall@10", 0.0),
                            scores.get("ndcg@10", 0.0),
                            scores.get("mrr@10", 0.0),
                            scores.get(candidate_hit_metric, 0.0),
                        ]
                    )
                print(f"Wrote metrics to {args.metrics_csv}")
        if temp_eval_dir is not None:
            temp_eval_dir.cleanup()
        return

    if not args.query:
        if temp_eval_dir is not None:
            temp_eval_dir.cleanup()
        raise SystemExit("Provide --query, or --doc-id, or --qrels (optionally with --queries).")

    hits = retrieve_claims(
        args.query,
        limit=args.limit,
        shard=args.retrieve_shard,
        retrieval_mode=args.retrieval_mode,
        hybrid_alpha=args.hybrid_alpha,
    )
    reranked = rerank_hits(
        hits,
        args.query,
        shard=args.rerank_shard,
        rerank_k=args.rerank_k,
        rerank_source=args.rerank_source,
    )
    if args.debug_compare_vectors and args.rerank_source == "weaviate":
        debug_compare_lmdb_vs_weaviate_vectors(
            reranked[: max(1, args.rerank_k)],
            args.rerank_shard,
            sample_n=args.debug_compare_n,
            max_abs_tol=args.debug_compare_max_abs_tol,
        )
    print(
        f"Retrieved {len(hits)} hits; reranked top {args.rerank_k} "
        f"using mode={args.retrieval_mode} alpha={args.hybrid_alpha:.3f} "
        f"retrieve_shard={args.retrieve_shard} rerank_shard={args.rerank_shard} "
        f"rerank_source={args.rerank_source}"
    )
    _print_hits(reranked, k=10)
    if args.save_before_csv:
        _write_hits_csv(hits, args.save_before_csv, k=max(1, args.save_k))
        print(f"Saved pre-rerank results to {args.save_before_csv}")
    if args.save_after_csv:
        _write_hits_csv(reranked, args.save_after_csv, k=max(1, args.save_k))
        print(f"Saved reranked results to {args.save_after_csv}")
    if temp_eval_dir is not None:
        temp_eval_dir.cleanup()


if __name__ == "__main__":
    main()
